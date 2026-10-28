"""助手导出：仅在响应内存中还原，Word/Excel 不落盘。"""
import io
import re
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from sqlalchemy import select

from app.assistant import service
from app.audit import service as audit
from app.db.models import DocNode, TaskFile
from app.export import report_docx
from app.security import crypto
from app.security.errors import AppError

HEADER = '商衡SI 智能助手 · 内容由 AI 生成，仅供参考'


def _cells(line):
    return [value.strip().replace('\\|', '|') for value in re.split(r'(?<!\\)\|', line.strip().strip('|'))]


def blocks(text):
    """只解析标题、段落、一层列表、粗体与表格；链接和 HTML 仍为文字。"""
    lines, i = text.splitlines(), 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if '|' in line and i + 1 < len(lines) and all(re.fullmatch(r':?-{3,}:?', c) for c in _cells(lines[i+1])):
            headers, rows = _cells(line), []
            i += 2
            while i < len(lines) and '|' in lines[i] and lines[i].strip():
                cells = _cells(lines[i])
                rows.append((cells + [''] * len(headers))[:len(headers)])
                i += 1
            yield 'table', (headers, rows)
            continue
        heading = re.match(r'^#{1,6}\s+(.+)', line)
        listing = re.match(r'^(?:[-*+] |\d+[.)] )(.+)', line)
        if heading:
            yield 'heading', heading.group(1)
        elif listing:
            yield 'list', listing.group(1)
        elif re.fullmatch(r'-{3,}|\*{3,}', line):
            yield 'paragraph', '────────'
        else:
            yield 'paragraph', line
        i += 1


def _inline(paragraph, text):
    for index, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if part:
            report_docx._font(paragraph.add_run(report_docx._clean(part)), 10.5, index % 2 == 1)


def _document():
    doc = report_docx._new_document()
    # 沿用已有报告字体、页面与边距，移除审核专用标题。
    first = doc.paragraphs[0]._element
    first.getparent().remove(first)
    report_docx._font(doc.sections[0].header.paragraphs[0].add_run(HEADER), 9)
    return doc


def _markdown_doc(doc, content):
    for kind, value in blocks(content):
        if kind == 'table':
            headers, rows = value
            report_docx._table(doc, headers, rows, [26.7 / len(headers)] * len(headers))
        elif kind == 'heading':
            report_docx._heading(doc, value)
        else:
            p = doc.add_paragraph()
            _inline(p, ('• ' if kind == 'list' else '') + value)


def _translation_doc(db, task, row, view, doc, variant):
    file_ids = row.params.get('file_ids', [])
    nodes = db.scalars(select(DocNode).join(TaskFile).where(DocNode.task_id == task.id,
                        TaskFile.id.in_(file_ids), DocNode.part == 'body')
                        .order_by(TaskFile.seq, DocNode.ord)).all()
    translated = view.get('translation', {})
    if variant == 'bilingual':
        rows = [(crypto.decrypt_text(n.text_enc) if n.text_enc else '原文已按保留期限删除',
                 translated.get(n.loc, '未完成翻译')) for n in nodes]
        report_docx._table(doc, ('原文', '译文'), rows, (13.35, 13.35))
        return
    if not translated:
        _markdown_doc(doc, view['content'])
        return
    table, table_group, table_row = None, None, None
    for node in nodes:
        value = translated.get(node.loc, '未完成翻译')
        match = re.match(r'^(.*)-R(\d+)-C(\d+)$', node.loc) if node.kind == 'table_cell' else None
        if node.kind == 'sheet_cell':
            from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
            group, _, ref = node.loc.rpartition('-')
            try:
                col, number = coordinate_from_string(ref)
                match = (group, number, column_index_from_string(col))
            except ValueError:
                match = None
        if match:
            group, number, col = (match.group(1), int(match.group(2)), int(match.group(3))) if hasattr(match, 'group') else match
            # 稀疏表格只生成实际出现的行列，避免异常坐标造成巨大空表。
            if table is None or group != table_group:
                table = doc.add_table(rows=0, cols=0); table.style = 'Table Grid'
                table_group, table_row = group, None
                columns = sorted({int(m.group(3)) for n in nodes if
                     (m := re.match(r'^(.*)-R(\d+)-C(\d+)$', n.loc)) and m.group(1) == group}) if node.kind == 'table_cell' else sorted({
                     column_index_from_string(coordinate_from_string(n.cell_ref)[0]) for n in nodes
                     if n.kind == 'sheet_cell' and n.loc.rpartition('-')[0] == group and n.cell_ref})
                if col not in columns:
                    columns.append(col); columns.sort()
                from docx.shared import Cm
                for _ in columns:
                    table.add_column(Cm(26.7 / len(columns)))
            if table_row != number:
                cells = table.add_row().cells; table_row = number
            _inline(cells[columns.index(col)].paragraphs[0], value)
        else:
            table, table_group, table_row = None, None, None
            if node.kind == 'heading':
                report_docx._heading(doc, value)
            else:
                _inline(doc.add_paragraph(), value)


def _comparison_rows(comparison):
    suppliers = comparison['suppliers']
    # 每项每家一行，便于数值格式保留与用户筛选。
    headers = ['项目', '供应方', '币种', '单价', '金额', '差额百分比']
    rows = []
    names = {item['file']: item['name'] for item in suppliers}
    for item in comparison['rows']:
        for quote in item['quotes']:
            rows.append([item['item'], names.get(quote['file'], ''), quote.get('currency', ''),
                         Decimal(quote['unit_price']) if quote.get('unit_price') is not None else None,
                         Decimal(quote['amount']) if quote.get('amount') is not None else None,
                         Decimal(quote['difference_pct']) / 100 if quote.get('difference_pct') is not None else None])
    for key, total in comparison['totals'].items():
        rows.append(['合计', names.get(int(key), ''), total.get('currency', ''), None,
                     Decimal(total['amount']) if total.get('amount') is not None else total.get('status', '不完整'), None])
    return headers, rows


def _xlsx(view):
    comparison = view.get('compare')
    if comparison:
        headers, rows = _comparison_rows(comparison)
    else:
        first = next((value for kind, value in blocks(view['content']) if kind == 'table'), None)
        if first is None:
            raise AppError(409, 'NO_TABLE', '回复中没有可导出的表格')
        headers, rows = first
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '智能助手'
    for values in [headers, *rows]:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.value = report_docx._clean(cell.value) if cell.value else ''
                cell.data_type = 's'  # 模型/原文中的公式、超链接都作为文字导出。
    if comparison:
        for row in sheet.iter_rows(min_row=2):
            for cell in row[3:5]:
                if cell.data_type == 'n':
                    cell.number_format = '#,##0.00'
            row[5].number_format = '0.0%'
        sheet.append(['价格未发送给模型'])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{sheet.cell(len(rows)+1, len(headers)).coordinate}'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    output = io.BytesIO(); workbook.save(output)
    return output.getvalue()


def generate(db, task, mid, format, variant, user, ip):
    row = service.owned_message(db, task, mid)
    if row.role != 'assistant' or row.status not in ('done', 'stopped'):
        raise AppError(409, 'ASSIST_EXPORT_UNAVAILABLE', '回复完成或停止后才能导出')
    view = service.message_out(db, task, row)
    if row.template == 'reply':
        view['content'] = view.get('draft', '')
        if not view['content'].strip():
            raise AppError(409, 'ASSIST_EXPORT_UNAVAILABLE', '未取得有效草稿，请重试')
    if format == 'xlsx':
        data = _xlsx(view)
        media = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        doc = _document()
        if row.template == 'translate':
            _translation_doc(db, task, row, view, doc, variant)
        elif view.get('compare'):
            headers, rows = _comparison_rows(view['compare'])
            rows = [values[:5] + [f'{values[5] * 100:.1f}%' if values[5] is not None else '—'] for values in rows]
            report_docx._table(doc, headers, rows, [26.7 / len(headers)] * len(headers))
            doc.add_paragraph('价格未发送给模型')
        else:
            _markdown_doc(doc, view['content'])
        output = io.BytesIO(); doc.save(output); data = output.getvalue()
        media = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    title = re.sub(r'[\\/:*?"<>|\x00-\x1f\x7f]', '_', task.title[:20])
    filename = f'智能助手-{title}-{datetime.now(ZoneInfo("Asia/Shanghai")):%Y%m%d%H%M}.{format}'
    audit.record(db, 'assistant_export', user_id=user.id, target_type='task', target_id=str(task.id),
                 detail={'message_id': row.id, 'format': format}, ip=ip)
    db.commit()
    return data, media, filename
