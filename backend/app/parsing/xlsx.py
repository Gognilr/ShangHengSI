"""XLSX 解析（M3-02，SPEC 8.4，设计 docs/design/M3.md 第三节，open_questions Q-62）。

openpyxl 读两次：data_only=True 取缓存值，data_only=False 判断公式。每个非空单元格一个节点。
解析失败只抛出 ParseError，消息是固定的中文原因，不含原文和异常消息。日志不写单元格内容。
"""
import io
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

from app.parsing.model import MAX_XLSX_CELLS, ParsedFile, ParsedNode, ParseError

PARSER_VERSION = "xlsx-1"

MSG_CORRUPT = "XLSX 文件无法读取，可能已损坏或已加密"
MSG_TOO_MANY_CELLS = "单个 XLSX 的非空单元格超过 20000 个"
MSG_EMPTY = "未能从文件中读取到内容"
WARN_FORMULA = "含未计算的公式，相关金额无法核验"
WARN_HIDDEN_ROWCOL = "含隐藏的行或列，已按正常内容审核"
WARN_HIDDEN_SHEET = "隐藏工作表未参与审核"

_CURRENCY_SYMBOLS = "¥￥$€£"
_BRACKET_CURRENCY = re.compile(r"\[\$([^\]\-]*)(?:-[^\]]*)?\]")  # [$USD]、[$¥-804]、[$-409]
_QUOTED = re.compile(r'"[^"]*"')


def _percent(fmt: str) -> bool:
    return "%" in _QUOTED.sub("", fmt)


def _currency(fmt: str) -> str | None:
    """格式含货币符号或代码时返回它，否则 None。"""
    m = _BRACKET_CURRENCY.search(fmt)
    if m and m.group(1):
        return m.group(1)
    rest = _BRACKET_CURRENCY.sub("", fmt)
    for ch in rest:
        if ch in _CURRENCY_SYMBOLS:
            return ch
    return None


def _to_decimal(value) -> Decimal:
    # float 用 repr（最短往返表示）转换，避免 Decimal(float) 带出二进制展开
    return value if isinstance(value, Decimal) else Decimal(repr(value) if isinstance(value, float) else str(value))


def _decimal_text(d: Decimal) -> str:
    """不含科学计数法、去掉尾随 0 的十进制文字。"""
    text = format(d, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return "0" if text in ("-0", "") else text


def cell_text(value, fmt: str) -> str:
    """单元格值转文字（Q-62 ①）。空白返回空字符串。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float, Decimal)):
        try:
            d = _to_decimal(value)
        except InvalidOperation:
            return str(value)
        if not d.is_finite():
            return str(value)
        if _percent(fmt):
            return _decimal_text(d * 100) + "%"
        text = _decimal_text(d)
        cur = _currency(fmt)
        return f"{cur} {text}" if cur else text
    return str(value).strip()


def _load(data: bytes, data_only: bool):
    try:
        return load_workbook(io.BytesIO(data), data_only=data_only)
    except Exception:  # noqa: BLE001 —— 任何读取失败都只给固定提示，不带异常消息（可能含原文）
        raise ParseError(MSG_CORRUPT) from None


def parse_xlsx(data: bytes, file_seq: int) -> ParsedFile:
    wb_val = _load(data, True)
    wb_formula = _load(data, False)

    has_uncached = has_hidden_rowcol = has_hidden_sheet = False
    nodes: list[ParsedNode] = []

    for idx, ws in enumerate(wb_val.worksheets, start=1):
        wf = wb_formula[ws.title]
        hidden = ws.sheet_state != "visible"
        has_hidden_sheet |= hidden
        has_hidden_rowcol |= any(d.hidden for d in ws.row_dimensions.values())
        has_hidden_rowcol |= any(d.hidden for d in ws.column_dimensions.values())
        # ws._cells 是只含已写入单元格的稀疏表；iter_rows 会遍历整个声明范围，声明很大的空表会很慢
        for (row, col), cell in sorted(ws._cells.items()):
            formula_cell = wf._cells.get((row, col))
            if formula_cell is not None and formula_cell.data_type == "f" and cell.value is None:
                has_uncached = True
                continue
            text = cell_text(cell.value, cell.number_format)
            if not text:
                continue
            if len(nodes) >= MAX_XLSX_CELLS:
                raise ParseError(MSG_TOO_MANY_CELLS)
            ref = f"{get_column_letter(col)}{row}"
            nodes.append(
                ParsedNode(
                    loc=f"F{file_seq}-S{idx}-{ref}",
                    ord=len(nodes) + 1,
                    kind="sheet_cell",
                    part="hidden_sheet" if hidden else "body",
                    text=text,
                    sheet=ws.title,
                    cell_ref=ref,
                )
            )

    warnings: list[str] = []
    if has_uncached:
        warnings.append(WARN_FORMULA)
    if has_hidden_rowcol:
        warnings.append(WARN_HIDDEN_ROWCOL)
    if has_hidden_sheet:
        warnings.append(WARN_HIDDEN_SHEET)
    if not nodes:
        raise ParseError(MSG_EMPTY)
    return ParsedFile(nodes=nodes, page_count=None, has_revisions=False, warnings=warnings)
