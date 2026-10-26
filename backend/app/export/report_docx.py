"""SPEC 14 Word 报告；只读轮次快照、校验结果和节点位置元数据。"""
import io
import re
from datetime import datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy import select

from app.checks.service import list_findings
from app.db.models import ReviewRun, Task, TaskFile, User
from app.review.prompt import DOC_TYPE_LABELS, SOURCE_LABELS, STANCE_LABELS
from app.review.service import run_detail

DECLARATION = "本报告由 AI 辅助生成，仅供内部参考，不构成法律意见。重大事项及高风险交易请由法务、合规人员或律师确认。"
SANCTIONS_NOTE = "未联网核查，基于模型自身知识，需合规人员确认最新制裁与出口管制规定"
SEVERITY = {"high": "高", "medium": "中", "low": "低"}
CATEGORY = {"wording": "表述", "logic": "逻辑", "legal": "法律", "compliance": "合规与制裁", "commercial": "商务"}
CITE = {"verified": "引用已校验", "unverified": "待核实", "missing": "缺失型"}
HANDLING = {"pending": "待处理", "adopted": "已采纳", "ignored": "已忽略"}
VERDICT = {"issue": "有问题", "ok": "未发现问题", "na": "不适用", "unable": "无法判断"}
CHECKS = dict(zip(("L-01", "L-02", "L-03", "L-04", "L-05", "L-06", "L-07", "L-08", "L-INFO"),
                  ("行金额", "合计", "税额", "大小写一致", "跨文件金额", "币种一致", "报价有效期", "交期与签约日", "提示")))
CHECKS.update({'L-09': '分期金额', 'L-10': '废止法律', 'L-11': '增值税税率'})
SANCTIONS = {"none": "无明显风险", "attention": "需注意", "high": "高风险（建议合规人员或律师确认）"}


def _clean(value) -> str:
    text = str(value) if value is not None and value != "" else "—"
    return "".join(c for c in text if c in "\t\n\r" or "\x20" <= c <= "\ud7ff" or "\ue000" <= c <= "\ufffd" or "\U00010000" <= c <= "\U0010ffff")


def _font(run, size, bold=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def _new_document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(6)
    title = doc.add_paragraph(style="Title")
    # 清除 python-docx 默认 Title 的主题色和底边线，报告标题只用黑色文字。
    title_style = doc.styles["Title"]
    for border in list(title_style.element.xpath("./w:pPr/w:pBdr")):
        border.getparent().remove(border)
    title_run = title.add_run("商务文件审核报告")
    _font(title_run, 16, True)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def _heading(doc, title):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8)
    _font(paragraph.add_run(title), 10.5, True)


def _table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for cell, text in zip(table.rows[0].cells, headers):
        _font(cell.paragraphs[0].add_run(_clean(text)), 9, True)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            _font(cell.paragraphs[0].add_run(_clean(value)), 9)
    # 列与单元格宽度同时设置，避免 Word 使用默认等宽布局。
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
    return table


def _files(db, task):
    files = list(db.scalars(select(TaskFile).where(TaskFile.task_id == task.id).order_by(TaskFile.seq)))
    main = next((file for file in files if file.is_main), None)
    return files, main


def _filename(main, task):
    name = main.original_name if main else task.title
    stem = name.rsplit(".", 1)[0] if "." in name else name
    stem = re.sub(r'[\\/:*?"<>|\x00-\x1f\x7f]', "_", stem) or "文件"
    return f"审核报告_{stem}_{datetime.now(ZoneInfo('Asia/Shanghai')):%Y%m%d}.docx"


def _basic(doc, task, files, main, user, run=None):
    _heading(doc, "一 基本信息")
    filenames = [main.original_name if main else task.title] + [f"附件：{file.original_name}" for file in files if not file.is_main]
    if run is None:
        version, when = "—", "—"
    else:
        version = f"内置 {run.builtin_req_version or '—'}；" + (
            f"公司补充要求 第 {run.company_req_version} 版" if run.company_req_version is not None else "公司补充要求：无")
        when = run.finished_at.astimezone(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M") if run.finished_at else "—"
    rows = [("文件名", "\n".join(filenames)),
            ("文件类型", DOC_TYPE_LABELS.get(run.main_doc_type if run else main.doc_type if main else None, "—")),
            ("立场", STANCE_LABELS.get(run.stance if run else task.stance, "—")),
            ("来源", SOURCE_LABELS.get(run.source if run else task.source, "—")),
            ("审核人", user.display_name), ("审核时间", when), ("模型", run.model_name if run else "—"),
            ("审核要求版本", version), ("轮次", f"第 {run.run_no} 轮" if run else "—")]
    _table(doc, ("项目", "内容"), rows, (3.5, 23.2))
    if run is None:
        doc.add_paragraph("未做 AI 审核")


def _local(doc, db, task):
    _heading(doc, "五 本地校验结果")
    items = list_findings(db, task)
    if not items:
        doc.add_paragraph("本地校验未发现问题")
    else:
        _table(doc, ("校验项", "严重程度", "位置", "说明"),
               [(CHECKS.get(row["check_code"], row["check_code"]), SEVERITY.get(row["severity"], row["severity"]),
                 row["location"], row["detail"]) for row in items], (3, 2, 5.5, 16.2))


def _finish(doc):
    _heading(doc, "七 固定声明")
    doc.add_paragraph(DECLARATION)
    data = io.BytesIO()
    doc.save(data)
    return data.getvalue()


def build_run_report(db, run: ReviewRun, user: User) -> tuple[str, bytes]:
    task = db.get(Task, run.task_id)
    files, main = _files(db, task)
    detail = run_detail(db, run)  # 复用结果页的位置、覆盖说明和还原，不读取节点原文。
    # run_detail 已对模型字段、NOTE 及要求文字做一次还原，引用快照也是原文；不能再次替换已还原值。
    text = _clean

    doc = _new_document()
    _basic(doc, task, files, main, user, run)
    if detail.template:
        state={'modified':'（范本已修改）','deleted':'（范本已删除）'}.get(detail.template['status'],'')
        doc.add_paragraph('本轮对照公司范本：'+text(detail.template['name'])+state)
    _heading(doc, "二 覆盖说明与限制事项")
    if detail.coverage:
        doc.add_paragraph(text(detail.coverage["summary"]))
        for note in detail.coverage["notes"]:
            doc.add_paragraph(text(note))
    else:
        doc.add_paragraph("本轮次没有覆盖说明")
    _heading(doc, "三 结论摘要")
    doc.add_paragraph("；".join(f"{label}问题 {sum(f.severity == code for f in detail.findings)} 条" for code, label in SEVERITY.items()))
    doc.add_paragraph("制裁结论：" + SANCTIONS.get(detail.sanctions.level, "—"))
    for reason in detail.sanctions.reasons:
        doc.add_paragraph(text(reason))
    doc.add_paragraph(SANCTIONS_NOTE)
    _heading(doc, "四 问题清单")
    rows = []
    for finding in detail.findings:
        problem = [text(finding.problem)]
        if finding.source=='template':problem.insert(0,'范本对照')
        problem += [label + text(value) for label, value in (("对我方的影响：", finding.impact), ("依据：", finding.basis)) if value]
        for ref in finding.law_refs:
            problem.append(f"依据：《{_clean(ref['title'])}》{_clean(ref.get('article_no') or '')}\n{_clean(ref['text'])}")
        if finding.law_unverified:
            problem.append('引用未经法规库核实')
        suggestion = ([text(finding.suggestion)] if finding.suggestion else [])
        if finding.suggested_text:
            suggestion.append("建议措辞：" + text(finding.suggested_text))
        handling = HANDLING.get(finding.handling_status, finding.handling_status)
        if finding.handling_note:
            handling += "\n" + _clean(finding.handling_note)
        rows.append((finding.seq, SEVERITY.get(finding.severity, finding.severity), CATEGORY.get(finding.category, finding.category),
                     finding.location or ("缺失项" if finding.missing else "—"), text(finding.quote), "\n".join(problem),
                     "\n".join(suggestion), CITE.get(finding.cite_status, finding.cite_status), handling))
    if rows:
        _table(doc, ("序号", "等级", "类别", "位置", "原文", "问题说明", "修改建议", "引用状态", "处理状态与备注"),
               rows, (1, 1, 1.5, 3, 4.5, 5.5, 5, 2.2, 3))
    else:
        doc.add_paragraph("本次没有列出问题条目。结果由 AI 生成，仅供参考，请以人工复核为准。")
    _local(doc, db, task)
    _heading(doc, "六 审核要求逐条结论")
    if detail.requirement_results:
        _table(doc, ("编号", "内容", "结论", "备注"),
               [(row.req_id, text(row.req_text), VERDICT.get(row.verdict, row.verdict), text(row.note))
                for row in detail.requirement_results], (2, 13, 2.7, 9))
    else:
        doc.add_paragraph("本轮次没有审核要求逐条结论")
    return _filename(main, task), _finish(doc)


def build_local_report(db, task: Task, user: User) -> tuple[str, bytes]:
    files, main = _files(db, task)
    doc = _new_document()
    _basic(doc, task, files, main, user)
    _local(doc, db, task)
    return _filename(main, task), _finish(doc)
