"""目录读取与差异计算；不访问网络、不记录文件内容。"""
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.models import Law
from app.parsing import legacy
from app.parsing.model import ParseError
from app.security.errors import AppError
from app.tasks import filecheck

MAX_FILE_ROWS = 5000
MAX_TOTAL_ROWS = 20000
STATUS_LABELS = dict(zip(['尚未生效', '有效', '已修改', '已废止'], ['not_yet', 'valid', 'amended', 'repealed']))
REQUIRED = ['标题', '公布日期', '法律法规分类']


class CatalogError(Exception):
    """只允许本模块固定文字或由受信列名构造的提示。"""


def iso(value):
    return value.isoformat() if isinstance(value, date) else value


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', value.strip()):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    return None


def _text(value):
    return str(value).strip() if value is not None else ''


def _workbook(data):
    try:
        if data.startswith(filecheck.OLE_MAGIC):
            filecheck.check_ole(io.BytesIO(data), ext='xls', name='目录表格')
            data = legacy.convert(data, 'xls')
        elif not data.startswith(b'PK'):
            raise CatalogError('不是官方导出的目录表格')
        filecheck.check_ooxml(io.BytesIO(data), 'xlsx', '目录表格')
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except AppError as exc:
        raise CatalogError(exc.message) from None
    except ParseError as exc:
        # converter 的提示固定；不透出第三方解析器的异常文字。
        safe = {legacy.MSG_UNAVAILABLE, legacy.MSG_TIMEOUT, legacy.MSG_FAILED}
        raise CatalogError(str(exc) if str(exc) in safe else legacy.MSG_FAILED) from None


def preview(db, files, declared_status):
    diff = {'added': [], 'changed': [], 'skipped': [], 'duplicates': 0}
    rows, seen = [], set()
    total = skipped = 0
    for seq, data in files:
        book = _workbook(data)
        try:
            sheet = book.worksheets[0]
            # 忽略工作表声明的维度，防止错误维度隐藏有效行或绕过行数限制。
            sheet.reset_dimensions()
            iterator = enumerate(sheet.iter_rows(values_only=True), 1)
            header = None
            for n, values in iterator:
                if n > 5:
                    break
                if any(_text(v) for v in values):
                    header = {_text(v): i for i, v in enumerate(values)}
                    break
            missing = [name for name in REQUIRED if header is None or name not in header]
            if missing:
                raise CatalogError(f'第 {seq} 个文件缺少列：'+'、'.join(missing))
            count = 0
            for n, values in iterator:
                count += 1
                total += 1
                if count > MAX_FILE_ROWS or total > MAX_TOTAL_ROWS:
                    raise CatalogError('目录表格行数超过限制（每份最多 5000 行，合计最多 20000 行）')
                def cell(name):
                    idx = header.get(name)
                    return values[idx] if idx is not None and idx < len(values) else None
                title, published = _text(cell('标题')), _date(cell('公布日期'))
                category, office = _text(cell('法律法规分类')), _text(cell('制定机关'))
                reason = ('标题为空' if not title else '标题超过 300 字' if len(title) > 300 else
                          '公布日期为空或不合法' if published is None else
                          '法律法规分类为空或超过 60 字' if not category or len(category) > 60 else
                          '制定机关超过 100 字' if len(office) > 100 else None)
                if reason:
                    skipped += 1
                    if len(diff['skipped']) < 200:
                        diff['skipped'].append({'file': seq, 'row': n, 'reason': reason})
                    continue
                key = (title, published)
                if key in seen:
                    diff['duplicates'] += 1
                    continue
                seen.add(key)
                rows.append(dict(title=title, published_on=published, category=category,
                                 effective_on=_date(cell('施行日期')),
                                 status=STATUS_LABELS.get(_text(cell('时效性')), declared_status)))
        finally:
            book.close()
    existing = {(r.title, r.published_on): r for r in db.scalars(select(Law))}
    unchanged = 0
    for row in rows:
        old = existing.get((row['title'], row['published_on']))
        if old is None:
            diff['added'].append({k: iso(v) for k, v in row.items()})
            continue
        changes = {k: [iso(getattr(old, k)), iso(row[k])] for k in ('status', 'category', 'effective_on')
                   if (k != 'effective_on' or row[k] is not None) and getattr(old, k) != row[k]}
        if changes:
            diff['changed'].append(dict(law_id=old.id, title=old.title,
                                        published_on=iso(old.published_on), changes=changes))
        else:
            unchanged += 1
    return diff, unchanged, skipped
