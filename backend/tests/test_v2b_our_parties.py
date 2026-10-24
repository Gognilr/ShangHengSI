"""V2-B-04：设计 4.4 测试先行。"""
import pytest
import io
from docx import Document
from tests.review_support import ready_task  # noqa: F401
from tests.test_redaction_api import alice  # noqa: F401
from app.redaction.engine import build_outbound, apply_party_changes
from app.db.models import Task
from app.redaction.engine import NodeIn, Term, compute, PartyOverride
from app.redaction.outbound import PartyInfo, party_description
from tests.test_admin_redaction import admin, env  # noqa: F401
from app.main import API_PREFIX

OUR = "澄川虚构信息技术有限公司"
OTHER = "青岚虚构石化股份有限公司"


def run(stance, terms=True):
    return compute([NodeIn("F1-P1", "甲方：" + OUR), NodeIn("F1-P2", "乙方：" + OTHER)], stance,
                   terms=[Term(OUR, "learned", "ORG", True)] if terms else [])


@pytest.mark.parametrize("stance,code,description", [
    ("purchase", "供应商_A", "我方=甲方（采购方）"), ("sales", "客户_A", "我方=甲方（销售方）"),
    ("other", "公司_A", "我方=甲方")])
def test_01_02_03_stance(stance, code, description):
    r = run(stance)
    assert next(p for p in r.parties if p.is_our_party).role_label == "甲方"
    assert r.redacted["F1-P2"] == "乙方：" + code
    infos = [PartyInfo(p.code, p.role_label, p.country, p.org_type, p.is_our_party, stance=stance) for p in r.parties]
    desc = party_description(infos)
    assert description in desc and code + "：乙方，中国" in desc


def test_04_parenthetical_role():
    r = compute([NodeIn("F1-P1", "甲方（采购方）：" + OUR)], "purchase")
    assert r.parties[0].role_label == "买方" and r.parties[0].is_our_party


def test_05_block_wins():
    r = compute([NodeIn("F1-P1", "甲方：" + OUR), NodeIn("F1-P2", OTHER + "负责交付")], "purchase",
                terms=[Term(name, "learned", "ORG", True) for name in (OUR, OTHER)])
    assert [p.key for p in r.parties if p.is_our_party] == [OUR.casefold()]


@pytest.mark.usefixtures("env")
def test_06_api(admin, make_user, login):
    c, csrf = admin
    url = API_PREFIX + "/admin/our-parties"
    h = {"X-CSRF-Token": csrf}
    row = c.post(url, json={"name": OUR}, headers=h)
    assert row.status_code == 201, row.text
    assert row.json()["is_our_party"] and row.json()["category"] == "ORG"
    assert len(c.get(url).json()["items"]) == 1
    bad = c.post(url, json={"name": "合同"}, headers=h)
    assert bad.status_code == 422 and bad.json()["code"] == "TERM_IN_PROMPT"
    assert c.delete(url + "/" + row.json()["id"], headers=h).status_code == 204
    make_user("ordinary")
    u, token = login("ordinary")
    assert u.get(url).status_code == 403
    assert u.post(url, json={"name": OUR}, headers={"X-CSRF-Token": token}).status_code == 403


def test_07_stance_recomputed():
    assert run("purchase").redacted["F1-P2"] == "乙方：供应商_A"
    assert run("sales").redacted["F1-P2"] == "乙方：客户_A"


@pytest.mark.usefixtures("env")
def test_preset_parse_and_stance_change(admin, alice, ready_task, db_sessionmaker):
    c, csrf = admin
    assert c.post(API_PREFIX + "/admin/our-parties", json={"name": OUR},
                  headers={"X-CSRF-Token": csrf}).status_code == 201
    user, _, _ = alice
    doc = Document()
    doc.add_paragraph("甲方：" + OUR)
    doc.add_paragraph("乙方：" + OTHER)
    data = io.BytesIO()
    doc.save(data)
    task = ready_task(user, files=[("虚构采购.docx", data.getvalue())], stance="purchase")
    with db_sessionmaker() as db:
        task = db.get(Task, task.id)
        assert "我方=甲方（采购方）；供应商_A：乙方，中国" in build_outbound(db, task).text
        task.stance = "sales"
        apply_party_changes(db, task, [])
        assert "我方=甲方（销售方）；客户_A：乙方，中国" in build_outbound(db, task).text


@pytest.mark.parametrize("labels", [("甲方", "乙方"), ("其他", "其他")])
def test_multiple_presets_ambiguous_not_selected(labels):
    nodes = [NodeIn("F1-P1", labels[0] + "：" + OUR), NodeIn("F1-P2", labels[1] + "：" + OTHER)]
    r = compute(nodes, "purchase", terms=[Term(n, "learned", "ORG", True) for n in (OUR, OTHER)])
    assert not any(p.is_our_party for p in r.parties)
