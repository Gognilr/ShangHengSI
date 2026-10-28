"""T-52：本人导出、真实值还原、保留期、管理员用途筛选。"""
import io
import uuid
from datetime import timedelta
from decimal import Decimal

import pytest
from docx import Document
from openpyxl import load_workbook
from sqlalchemy import select

from app.admin.settings_store import set_setting
from app.db.base import utcnow
from app.db.models import AssistantMessage, AuditLog, DocNode, Task, TaskFile, LlmCall
from tests.review_support import env  # noqa: F401
from tests.test_assistant_chat import conversation, BASE, REAL  # noqa: F401


def message(sm, tid, **values):
    with sm() as db:
        row = AssistantMessage(task_id=uuid.UUID(tid), role='assistant', status='done',
                               content_redacted=values.pop('content_redacted', '回复 [[ORG_1]]'), **values)
        db.add(row); db.commit()
        return row.id


def url(tid, mid, format='docx', variant='translation'):
    return f'{BASE}/{tid}/messages/{mid}/export?format={format}&variant={variant}'


def test_t52_1_6_status_and_owner(conversation, db_sessionmaker, make_user, login):
    c, h, tid, _ = conversation
    mid = message(db_sessionmaker, tid)
    other = make_user('export_other', role='admin')
    other_client, _ = login(other.username)
    assert other_client.get(url(tid, mid)).status_code == 404
    for status in ('queued', 'streaming', 'failed'):
        with db_sessionmaker() as db:
            db.get(AssistantMessage, mid).status = status
            db.commit()
        assert c.get(url(tid, mid)).status_code == 409
    with db_sessionmaker() as db:
        db.get(AssistantMessage, mid).status = 'stopped'
        db.commit()
    assert c.get(url(tid, mid)).status_code == 200


def test_t52_2_5_markdown_xlsx_and_no_table(conversation, db_sessionmaker):
    c, _, tid, _ = conversation
    mid = message(db_sessionmaker, tid, content_redacted='| 主体 | 要求 |\n|---|---|\n| [[ORG_1]] | 核对 |\n| =1+1 | 原样 |')
    response = c.get(url(tid, mid, 'xlsx'))
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert list(sheet.values) == [('主体', '要求'), (REAL, '核对'), ('=1+1', '原样')]
    assert sheet['A3'].data_type == 's'
    empty = message(db_sessionmaker, tid)
    response = c.get(url(tid, empty, 'xlsx'))
    assert response.status_code == 409 and response.json()['code'] == 'NO_TABLE'


def test_t52_3_comparison_numeric_xlsx(conversation, db_sessionmaker):
    from tests.test_assistant_compare import compare_input, data, MAPPING
    from app.assistant.compare import normalize
    c, _, tid, _ = conversation
    compare_input(db_sessionmaker, tid)
    mid = message(db_sessionmaker, tid, template='compare', extras={'compare': normalize(data(), MAPPING, [1, 2])})
    response = c.get(url(tid, mid, 'xlsx'))
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    numbers = [cell.value for row in sheet for cell in row if cell.data_type == 'n']
    assert 10 in numbers and 12.5 in numbers and 1000 in numbers and 1250 in numbers
    assert any(cell.data_type == 'n' and cell.value == 0.25 and '%' in cell.number_format for row in sheet for cell in row)


def test_t52_4_translation_bilingual_and_draft_only(conversation, db_sessionmaker):
    c, _, tid, _ = conversation
    with db_sessionmaker() as db:
        fid = str(db.scalar(select(TaskFile.id).where(TaskFile.task_id == uuid.UUID(tid))))
    mid = message(db_sessionmaker, tid, template='translate', params={'file_ids': [fid], 'target': '英文'},
                  extras={'translation': {'F1-P1': 'Translated [[ORG_1]]'}})
    response = c.get(url(tid, mid, variant='bilingual'))
    assert response.status_code == 200, response.text
    doc = Document(io.BytesIO(response.content))
    assert len(doc.tables[0].rows) == 2
    assert REAL in doc.tables[0].cell(1, 0).text and REAL in doc.tables[0].cell(1, 1).text
    assert doc.sections[0].header.paragraphs[0].text == '商衡SI 智能助手 · 内容由 AI 生成，仅供参考'
    draft = message(db_sessionmaker, tid, template='reply', content_redacted='说明不得导出\n---\n草稿 [[ORG_1]]')
    doc = Document(io.BytesIO(c.get(url(tid, draft)).content))
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert '说明不得导出' not in text and '草稿 '+REAL in text
    missing = message(db_sessionmaker, tid, template='reply', content_redacted='只有说明')
    assert c.get(url(tid, missing)).status_code == 409


def test_t52_7_8_original_then_result_retention(conversation, db_sessionmaker):
    from app.backup.retention import run_retention
    c, _, tid, _ = conversation
    mid = message(db_sessionmaker, tid)
    with db_sessionmaker() as db:
        task = db.get(Task, uuid.UUID(tid))
        task.last_activity_at = utcnow() - timedelta(days=3)
        for layer in ('original', 'result', 'outbound', 'audit', 'unfinished'):
            set_setting(db, f'retention.{layer}_days', 1 if layer == 'original' else 0, None)
        db.commit()
        assert run_retention(db)['tasks'] == 1
        assert db.scalar(select(DocNode).where(DocNode.task_id == task.id)).text_enc is None
    assert c.get(url(tid, mid)).status_code == 200
    response = c.post(f'{BASE}/{tid}/messages', headers=conversation[1], json={'message': '继续'})
    assert response.status_code == 409 and response.json()['code'] == 'SESSION_READONLY'
    with db_sessionmaker() as db:
        set_setting(db, 'retention.result_days', 2, None)
        db.commit()
        assert run_retention(db)['results'] == 1
        assert db.get(AssistantMessage, mid) is None


def test_t52_9_unfinished_preserves_confirmed_and_messages(conversation, db_sessionmaker):
    from app.backup.retention import run_retention
    c, _, tid, user = conversation
    with db_sessionmaker() as db:
        old = utcnow() - timedelta(days=5)
        db.get(Task, uuid.UUID(tid)).last_activity_at = old
        empty = Task(kind='assistant', owner_id=user.id, title='空会话', stance='other', source='incoming',
                     policy='ai', status='active', last_activity_at=old)
        has_message = Task(kind='assistant', owner_id=user.id, title='有消息', stance='other', source='incoming',
                           policy='ai', status='active', last_activity_at=old)
        db.add_all([empty, has_message]); db.flush()
        abandoned = Task(kind='assistant', owner_id=user.id, title='从未确认的旧失败解析', stance='other', source='incoming',
                         policy='ai', status='parse_failed', last_activity_at=old)
        db.add(abandoned); db.flush()
        db.add(TaskFile(task_id=abandoned.id, seq=1, is_main=True, original_name='虚构.docx', ext='docx',
                        size_bytes=1, sha256='2'*64, storage_key=uuid.uuid4().hex))
        db.add(AssistantMessage(task_id=has_message.id, role='user', status='done', content_redacted='问题'))
        for layer in ('original', 'result', 'outbound', 'audit', 'unfinished'):
            set_setting(db, f'retention.{layer}_days', 1 if layer == 'unfinished' else 0, None)
        db.commit()
        assert run_retention(db)['unfinished'] == 2
        assert db.get(Task, abandoned.id) is None
        assert db.get(Task, empty.id) is None and db.get(Task, has_message.id) is not None
        assert db.get(Task, uuid.UUID(tid)) is not None


def test_t52_10_assist_filter_and_count(conversation, db_sessionmaker, make_user, login):
    user = make_user('export_admin', role='admin')
    c, _ = login(user.username)
    with db_sessionmaker() as db:
        for purpose in ('assist', 'chat'):
            db.add(LlmCall(purpose=purpose, model_name='mock', status='sent'))
        db.commit()
    response = c.get('/api/v1/admin/outbound-logs?purpose=assist')
    assert response.status_code == 200
    assert response.json()['total'] == 1 and response.json()['items'][0]['purpose'] == 'assist'
    assert c.get('/api/v1/admin/system/status').json()['assistant_sessions'] == 1


def test_t52_11_export_audit_and_safe_filename(conversation, db_sessionmaker):
    c, h, tid, _ = conversation
    mid = message(db_sessionmaker, tid)
    response = c.get(url(tid, mid))
    assert response.status_code == 200 and "filename*=UTF-8''" in response.headers['content-disposition']
    with db_sessionmaker() as db:
        audits = db.scalars(select(AuditLog).where(AuditLog.action == 'assistant_export')).all()
        assert len(audits) == 1 and audits[0].detail == {'message_id': mid, 'format': 'docx'}
        assert REAL not in str(audits[0].detail) and '敏感会话标题' not in str(audits[0].detail)
        assert len(db.scalars(select(AuditLog).where(AuditLog.action == 'assistant_session_create')).all()) == 1
    assert c.delete(f'{BASE}/{tid}', headers=h).status_code == 204
    with db_sessionmaker() as db:
        assert len(db.scalars(select(AuditLog).where(AuditLog.action == 'assistant_session_delete')).all()) == 1


def test_translation_word_tables_and_original_deleted(conversation, db_sessionmaker):
    from app.security import crypto
    c, _, tid, _ = conversation
    with db_sessionmaker() as db:
        file = db.scalar(select(TaskFile).where(TaskFile.task_id == uuid.UUID(tid)))
        db.add_all(DocNode(task_id=file.task_id, file_id=file.id, loc=f'F1-T1-R{r}-C{col}', ord=1+r*2+col,
                           kind='table_cell', part='body', text_enc=crypto.encrypt_text('虚构单元格'),
                           redacted_text='[[ORG_1]]') for r in (1, 2) for col in (1, 2))
        fid = str(file.id)
        db.commit()
    translated = {'F1-P1': '段落'} | {f'F1-T1-R{r}-C{col}': f'译文{r}-{col} [[ORG_1]]' for r in (1, 2) for col in (1, 2)}
    mid = message(db_sessionmaker, tid, template='translate', params={'file_ids': [fid], 'target': '英文'},
                  extras={'translation': translated})
    doc = Document(io.BytesIO(c.get(url(tid, mid)).content))
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 2 and len(doc.tables[0].columns) == 2
    assert doc.tables[0].cell(1, 1).text == '译文2-2 '+REAL
    with db_sessionmaker() as db:
        for node in db.scalars(select(DocNode).where(DocNode.task_id == uuid.UUID(tid))):
            node.text_enc = None
        db.commit()
    doc = Document(io.BytesIO(c.get(url(tid, mid, variant='bilingual')).content))
    assert len(doc.tables[0].rows) == 6 and doc.tables[0].cell(1, 0).text == '原文已按保留期限删除'
