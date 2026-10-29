"""V2-B-06：实体识别接入。对拍、冻结留出集与参考质量三项只在更换模型时运行（设置 NER_MODEL_EVAL=1）。"""
import importlib.util
import json
import os
from pathlib import Path
from unittest.mock import patch

import pytest

FIXTURES = Path(__file__).parent / 'fixtures/ner'
model_eval = pytest.mark.skipif(not os.environ.get('NER_MODEL_EVAL'), reason='更换模型时设置 NER_MODEL_EVAL=1 运行')


def corpus():
    # V2Q-48：对拍须固定06a原始输入，不能随当前规则的遮盖范围改变。
    import hashlib
    data = (FIXTURES / 'reference_corpus.json').read_bytes()
    assert hashlib.sha256(data).hexdigest() == '63e2892c2ce663229d3ffaa5d95b34cfbd764f2d08f852b472a8e6e63fcb9ab1'
    return json.loads(data)['records']


@model_eval
def test_reference_parity():
    from app.redaction import ner
    reference = json.loads((FIXTURES / 'reference_predictions.json').read_text(encoding='utf-8-sig'))
    records = corpus()
    assert len(records) == len(reference) == 391
    largest = 0.0
    for index, record in enumerate(records):
        actual = {(e.start, e.end, e.label): e.score for e in ner.raw_recognize(record['input']) if e.score >= 0.3}
        expected = {(e['start'], e['end'], e['label']): e['score'] for e in reference[record['id']]}
        assert actual.keys() == expected.keys(), record['id']
        delta = max((abs(actual[k] - expected[k]) for k in actual), default=0)
        largest = max(largest, delta)
        assert delta <= 1e-4, (record['id'], delta)
        if index % 50 == 0:
            print('parity nodes', index, flush=True)
    print('parity max score difference', largest)


@model_eval
def test_frozen_holdout():
    from app.redaction import ner
    samples = json.loads((FIXTURES / 'holdout/samples.json').read_text(encoding='utf-8-sig'))['samples']
    assert len(samples) >= 8
    total = hits = 0
    misses = []
    for sample in samples:
        text = sample['text']
        predictions = ner.recognize(text)
        for surface, label in sample['entities']:
            assert text.count(surface) == 1
            start = text.index(surface)
            total += 1
            hit = any((e.start, e.end, e.label) == (start, start + len(surface), label)
                      and e.score >= ner.MEDIUM for e in predictions)
            hits += hit
            if not hit:
                misses.append((sample['id'], surface, label))
        for surface in sample['must_keep']:
            start = text.index(surface)
            assert not any(e.score >= ner.HIGH and e.start < start + len(surface) and start < e.end
                           for e in predictions), sample['id']
    print('holdout hits/total', hits, total, 'misses', misses)
    assert total >= 40
    assert hits / total >= 0.9


@model_eval
def test_reference_quality_and_filtered_false_positives():
    from app.redaction import ner
    reference = json.loads((FIXTURES / 'reference_predictions.json').read_text(encoding='utf-8-sig'))
    false = total = hits = 0
    for record in corpus():
        predictions = [ner.Entity(**p) for p in reference[record['id']]
                       if p['score'] >= ner.MEDIUM and ner.keep_entity(record['input'], ner.Entity(**p))]
        expected = {tuple(e) for e in record['expected']}
        false += sum((e.start, e.end, e.label) not in expected for e in predictions)
        if record['origin'] == 'narrative':
            for a, b, label in expected:
                if label in ('PERSON', 'ORG'):
                    total += 1
                    hits += any((e.start, e.end, e.label) == (a, b, label) for e in predictions)
        for a, b in record['keep']:
            assert not any(e.score >= ner.HIGH and e.start < b and a < e.end for e in predictions), record['id']
    print('narrative hits/total', hits, total, 'filtered false positives', false)
    assert hits / total >= 0.9
    assert false <= 20


def test_engine_masking_and_note(monkeypatch):
    from app.redaction import ner
    from app.redaction.engine import NodeIn, compute
    seen = []

    def predict(text):
        seen.append(text)
        return [ner.Entity(text.index('沈柏宁'), text.index('沈柏宁') + 3, 'PERSON', 0.99)]

    monkeypatch.setattr(ner, 'recognize', predict)
    text = '电话：13800138000，沈柏宁复核。'
    result = compute([NodeIn('NOTE', text)], 'other')
    assert '13800138000' not in seen[0]
    assert len(seen[0]) == len(text)
    assert any(i.detected_by == 'ner' and i.category == 'PERSON' for i in result.items)
    assert '沈柏宁' not in result.redacted['NOTE']


def test_high_confidence_name_replaced_in_every_node(monkeypatch):
    """V2-B-09：模型只在一个节点给出高置信时，其他节点（含 NOTE、在前的节点）的同名文字也要替换，
    否则对照表中有该真实值，出口复检会拦截发送；中置信疑似与已撤销的人名不受影响。"""
    import re
    from app.redaction import ner
    from app.redaction.engine import NodeIn, compute
    from app.redaction.recheck import recheck
    from app.redaction.restore import restore

    def predict(text):
        if '[[命中]]' not in text:
            return [ner.Entity(m.start(), m.end(), 'ORG', 0.7)
                    for m in re.finditer('Northern Fictive Laboratory', text)]
        return [ner.Entity(m.start(), m.end(), label, 0.99)
                for name, label in (('Northern Fictive Laboratory', 'ORG'), ('钱虚明', 'PERSON'), ('周示例', 'PERSON'))
                for m in re.finditer(name, text)]

    monkeypatch.setattr(ner, 'recognize', predict)
    nodes = [NodeIn('F1-P1', '甲方：青岚虚构石化股份有限公司'),
             NodeIn('F1-P2', '检测报告由 Northern Fictive Laboratory 出具，钱虚明负责，周示例抄送。'),
             NodeIn('F1-P3', '[[命中]] 据 Northern Fictive Laboratory 的结论，钱虚明、周示例签字。'),
             NodeIn('NOTE', '请钱虚明再核对 Northern Fictive Laboratory 的报告。')]
    result = compute(nodes, 'purchase', exclusions=['周示例'])
    outbound = '\n'.join(result.redacted.values())
    for name in ('Northern Fictive Laboratory', '钱虚明'):
        assert name not in outbound
    assert outbound.count('周示例') == 2  # 已撤销：全文不替换
    assert not result.suspects  # 中置信区间已被替换，不再重复列为疑似遗漏
    by_loc = {}
    for i in result.items:
        by_loc.setdefault(i.loc, set()).add((i.category, i.detected_by))
    assert {('ORG', 'ner'), ('PERSON', 'ner')} <= by_loc['F1-P2'] and {('ORG', 'ner'), ('PERSON', 'ner')} <= by_loc['NOTE']
    assert len({i.placeholder for i in result.items if i.category == 'ORG' and i.detected_by == 'ner'}) == 1
    assert not [r for r in recheck(outbound, real_values=result.mapping.values(), terms=[],
                                   placeholders=[k.split('#', 1)[0] for k in result.mapping]) if r.category == 'MAP']
    for n in nodes:
        assert restore(result.redacted[n.loc], result.mapping) == n.text


def test_chat_confidence_and_rule_precedence(monkeypatch):
    from app.redaction import ner
    from app.chat.service import _redact

    def predict(text):
        assert '13800138000' not in text
        return [ner.Entity(text.index('沈柏宁'), text.index('沈柏宁') + 3, 'PERSON', 0.99),
                ner.Entity(text.index('陆清禾'), text.index('陆清禾') + 3, 'PERSON', 0.7)]

    monkeypatch.setattr(ner, 'recognize', predict)
    result = _redact('电话：13800138000，沈柏宁复核陆清禾。', {}, [])
    assert '沈柏宁' not in result.text
    assert any(s['text'] == '陆清禾' and s['reason'] == 'ner' for s in result.suspects)


@pytest.mark.parametrize('surface,label', [('甲方', 'ORG'), ('Mr.', 'PERSON'), ('1', 'PERSON'),
                                          ('12 / 34-5', 'ADDR'), (' / ', 'ADDR')])
def test_filter(surface, label):
    from app.redaction import ner
    assert not ner.keep_entity(surface, ner.Entity(0, len(surface), label, 0.99))


@pytest.mark.parametrize('score', [0.7, 0.99])
def test_revoked_person_not_reintroduced_by_model(monkeypatch, score):
    from app.redaction import ner
    from app.redaction.engine import NodeIn, compute

    def predict(text):
        return [ner.Entity(m.start(), m.end(), 'PERSON', score)
                for m in __import__('re').finditer('周示例', text)]

    monkeypatch.setattr(ner, 'recognize', predict)
    result = compute([NodeIn('F1-P1', '联系人：周示例'), NodeIn('NOTE', '周示例复核。')],
                     'other', exclusions=['周示例'])
    assert not result.items
    assert not result.suspects
    assert not result.mapping


@pytest.mark.parametrize('surface,label', [('青岛港', 'ORG'), ('天津港', 'ADDR'),
    ('FOB Qingdao Port', 'ORG'), ('CIF Example', 'ORG'), ('HS 27131200', 'ADDR'),
    ('XC-132A', 'ADDR')])
def test_preserved_ports_terms_and_codes(surface, label):
    from app.redaction import ner
    assert not ner.keep_entity(surface, ner.Entity(0, len(surface), label, 0.99))


@pytest.mark.parametrize('surface,label', [('沈澄港', 'PERSON'), ('Anne Port', 'PERSON'),
                                          ('123 Port', 'ADDR')])
def test_port_filter_does_not_remove_person_or_street_address(surface, label):
    from app.redaction import ner
    assert ner.keep_entity(surface, ner.Entity(0, len(surface), label, 0.99))


@pytest.mark.parametrize('entry', ['app', 'worker'])
@pytest.mark.parametrize('changed', [False, True])
def test_startup_rejects_corrupt_or_missing_files(monkeypatch, tmp_path, entry, changed):
    import asyncio
    import hashlib
    from app.redaction import ner
    from app import main, worker
    monkeypatch.setattr(ner, 'MODEL_DIR', tmp_path)
    monkeypatch.setattr(main, 'check_master_key', lambda: None)
    monkeypatch.setattr(worker, 'check_master_key', lambda: None)
    digest = hashlib.sha256(b'original').hexdigest()
    for name in ner.FILES:
        (tmp_path / name).write_bytes(b'original')
    (tmp_path / 'MANIFEST.sha256').write_text(''.join(f'{digest}  {name}\n' for name in ner.FILES))
    if changed:
        (tmp_path / 'model.onnx').write_bytes(b'PRIVATE_FILE_CONTENT')
    else:
        (tmp_path / 'tokenizer.json').unlink()

    async def start_app():
        async with main.lifespan(main.app):
            pytest.fail('文件异常时不得进入服务期')

    with pytest.raises(RuntimeError, match='模型文件缺失或 SHA-256') as error:
        asyncio.run(start_app()) if entry == 'app' else worker.main()
    assert 'PRIVATE_FILE_CONTENT' not in str(error.value)


def test_model_load_failure_is_safe(monkeypatch):
    from app.redaction import ner
    monkeypatch.setattr(ner, '_model', None)
    monkeypatch.setattr(ner, 'check', lambda **kw: None)

    def failed():
        raise ValueError('PRIVATE_MODEL_CONTENT')

    monkeypatch.setattr(ner, '_Recognizer', failed)
    with pytest.raises(RuntimeError, match='模型加载失败') as error:
        ner._get_model()
    assert 'PRIVATE_MODEL_CONTENT' not in str(error.value)


def test_offline_body_note_and_question_with_real_model():
    import io
    from docx import Document
    from app.parsing.docx import parse_docx
    from app.redaction.engine import NodeIn, compute
    from app.chat.service import _redact
    text = '经与王建国沟通，沈柏宁复核了江岚研究院的记录。'
    document = Document()
    document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    parsed = parse_docx(stream.getvalue(), 1)
    nodes = [NodeIn(n.loc, n.text, n.clause_no) for n in parsed.nodes if n.part == 'body']
    result = compute([*nodes, NodeIn('NOTE', text)], 'other')
    for loc in ('F1-P1', 'NOTE'):
        assert any(i.loc == loc and i.detected_by == 'ner' for i in result.items)
        assert '沈柏宁' not in result.redacted[loc] or any(s.loc == loc for s in result.suspects)
    question = _redact('经与王建国沟通……', {}, [])
    assert '王建国' not in question.text or any(s['text'] == '王建国' for s in question.suspects)


def test_real_model_promotion_keeps_remaining_decisions_required():
    from app.chat.service import _redact
    from app.security.errors import AppError
    name = 'Meridian Fictive Inspection'
    message = name + ' 的检验报告能用吗？ Northern Fictive Laboratory 的报告呢？'
    preview = _redact(message, {}, [])
    assert preview.new_entries == {'[公司_Q1]': 'Northern Fictive Laboratory'}
    assert preview.suspects == [{'start': 0, 'end': len(name), 'text': name, 'reason': 'capitalized'}]
    with pytest.raises(AppError) as pending:
        _redact(message, {}, [], [])
    assert pending.value.code == 'CHAT_SUSPECTS_PENDING'
    decision = {'start': 0, 'end': len(name), 'decision': 'keep'}
    sent = _redact(message, {}, [], [decision])
    assert 'Northern Fictive Laboratory' not in sent.text
    assert sent.new_entries == preview.new_entries
    with pytest.raises(AppError) as duplicate:
        _redact(message, {}, [], [decision, decision])
    assert duplicate.value.status == 422
