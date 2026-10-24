"""V2-B-07：确认页十类手动补充、金额规范化与可撤销元数据。"""
import io

import pytest
from docx import Document

from app.redaction import ner
from app.redaction.engine import NodeIn, UserTerm, compute
from tests.review_support import env, ready_task  # noqa: F401
from tests.test_user_note_redaction import alice, task_url  # noqa: F401


@pytest.fixture(autouse=True)
def no_model(monkeypatch):
    monkeypatch.setattr(ner, 'recognize', lambda text: [])


@pytest.mark.parametrize('category,surface,placeholder', [
    ('ORG','虚构主体','公司_A'), ('PERSON','虚构人员','人员_1'), ('IM','fiction.account','[联系方式]'),
    ('ACCOUNT','765432109876','[账号_1]'), ('IDNO','AB765432','[证件号]'), ('REFNO','FICT-782','编号_1'),
    ('ADDR','中国虚构园区','[地址：中国]'), ('AMOUNT','1234.50','[金额_1]'),
    ('PROJECT','虚构试验项目','项目_1'), ('CUSTOM','虚构词条','[敏感词_1]')])
def test_ten_manual_categories(category, surface, placeholder):
    result = compute([NodeIn('F1-P1', surface)], 'other', user_terms=[UserTerm(surface, category)])
    assert result.redacted['F1-P1'] == placeholder
    assert result.items[0].detected_by == 'user'
    assert surface in result.mapping.values()


def test_manual_amount_reuses_normalized_automatic_code():
    result = compute([NodeIn('F1-P1','金额：1000.00'), NodeIn('F1-P2','参考1000')], 'other',
                     user_terms=[UserTerm('1000', 'AMOUNT')])
    assert [i.placeholder for i in result.items] == ['[金额_1]', '[金额_1]']
    assert result.redacted['F1-P1'] == '金额：[金额_1]'


def test_manual_amount_keeps_existing_currency_boundary():
    result = compute([NodeIn('F1-P1','人民币 1000.00'), NodeIn('F1-P2','参考1000')], 'other',
                     user_terms=[UserTerm('1000', 'AMOUNT')])
    assert len(result.items) == 2
    assert result.items[0].placeholder != result.items[1].placeholder


@pytest.mark.usefixtures('env')
@pytest.mark.parametrize('surface', ['不是数字', 'NaN', 'Infinity', '1,23'])
def test_invalid_manual_amount_rejected(alice, ready_task, surface):
    user, client, headers = alice
    document = Document()
    document.add_paragraph(surface)
    stream = io.BytesIO()
    document.save(stream)
    task = ready_task(user, files=[('虚构说明.docx',stream.getvalue())])
    response = client.post(task_url(task,'/redaction/manual'), headers=headers,
        json={'loc':'F1-P1','char_start':0,'char_end':len(surface),'category':'AMOUNT'})
    assert (response.status_code,response.json()['code']) == (422,'INVALID_AMOUNT')


@pytest.mark.usefixtures('env')
def test_revocability_is_authoritative_per_source(alice, ready_task, monkeypatch):
    user, client, headers = alice

    def recognize(text):
        start = text.find('沈柏宁')
        return [ner.Entity(start, start + 3, 'PERSON', 0.99)] if start >= 0 else []

    monkeypatch.setattr(ner, 'recognize', recognize)
    document = Document()
    for line in ('卖方：星河虚构进出口有限公司','联系人：周示例','沈柏宁复核。','虚构标记'):
        document.add_paragraph(line)
    stream = io.BytesIO()
    document.save(stream)
    task = ready_task(user, files=[('虚构说明.docx',stream.getvalue())])
    items = client.get(task_url(task,'/redaction')).json()['items']
    by_loc = {i['loc']: i for i in items}
    assert by_loc['F1-P1']['can_revoke'] is False
    assert by_loc['F1-P2']['can_revoke'] is True
    assert by_loc['F1-P3']['can_revoke'] is False
    manual = client.post(task_url(task,'/redaction/manual'), headers=headers,
        json={'loc':'F1-P4','char_start':0,'char_end':4,'category':'PROJECT'})
    assert manual.status_code == 200, manual.text
    target = next(i for i in manual.json()['items'] if i['loc']=='F1-P4')
    assert target['can_revoke'] is True
    response = client.post(task_url(task,'/redaction/revoke'), headers=headers,
        json={k:target[k] for k in ('loc','char_start','char_end')})
    assert response.status_code == 200
    assert not any(i['loc']=='F1-P4' for i in response.json()['items'])
