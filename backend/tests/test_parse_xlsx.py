"""XLSX 解析（M3-02）：manifest.yaml 中的预期事实、格式化、上限和损坏文件。"""
import io
from decimal import Decimal
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from app.parsing.model import MAX_XLSX_CELLS, ParseError
from app.parsing.xlsx import MSG_CORRUPT, MSG_EMPTY, MSG_TOO_MANY_CELLS, cell_text, parse_xlsx

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "xlsx"
MANIFEST = yaml.safe_load((FIXTURES / "manifest.yaml").read_text(encoding="utf-8"))["samples"]


def parse_sample(name: str):
    return parse_xlsx((FIXTURES / MANIFEST[name]["file"]).read_bytes(), 1)


def to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("name", list(MANIFEST))
def test_manifest(name):
    sample = MANIFEST[name]
    parsed = parse_sample(name)
    assert [(n.loc, n.text, n.part) for n in parsed.nodes] == [
        (e["loc"], e["text"], e["part"]) for e in sample["nodes"]
    ]
    assert parsed.warnings == sample["warnings"]
    assert parsed.page_count is None
    assert parsed.has_revisions is False
    assert parsed.ocr_pages == 0
    assert [n.ord for n in parsed.nodes] == list(range(1, len(parsed.nodes) + 1))
    assert all(n.kind == "sheet_cell" for n in parsed.nodes)
    locs = {n.loc for n in parsed.nodes}
    for loc in sample["no_node_locs"]:
        assert loc not in locs
    names = [s["name"] for s in sample["sheets"]]
    for n in parsed.nodes:
        assert n.loc == f"F1-S{names.index(n.sheet) + 1}-{n.cell_ref}"


@pytest.mark.parametrize(
    "value,fmt,expected",
    [
        (1250.0, "General", "1250"),
        (0.125, "General", "0.125"),
        (1e-7, "General", "0.0000001"),
        (1.5e10, "General", "15000000000"),
        (0.3, "0%", "30%"),
        (0.0125, "0.00%", "1.25%"),
        (1250, "[$USD] #,##0.00", "USD 1250"),
        (800000, '"¥"#,##0.00', "¥ 800000"),
        (5, "[$€-407]#,##0", "€ 5"),
        (5, "$#,##0", "$ 5"),
        (5, "[$-409]#,##0", "5"),
        (5, '"USD" 0', "5"),
        (-0.0, "General", "0"),
        (True, "General", "TRUE"),
        ("  abc  ", "General", "abc"),
        ("   ", "General", ""),
        (Decimal("2.50"), "General", "2.5"),
    ],
)
def test_cell_text(value, fmt, expected):
    assert cell_text(value, fmt) == expected


def test_no_float_artifacts_for_money():
    wb = Workbook()
    wb.active["A1"] = 1250.5
    wb.active["A2"] = 400000.0
    nodes = parse_xlsx(to_bytes(wb), 3).nodes
    assert [n.text for n in nodes] == ["1250.5", "400000"]
    assert nodes[0].loc == "F3-S1-A1"


def test_merged_cell_only_top_left():
    parsed = parse_sample("quotation")
    assert not any(n.cell_ref in ("B1", "G1", "B10", "E10") for n in parsed.nodes)


def test_cell_limit():
    wb = Workbook()
    ws = wb.active
    for i in range(1, MAX_XLSX_CELLS + 2):
        ws.cell(row=i, column=1, value="x")
    with pytest.raises(ParseError) as e:
        parse_xlsx(to_bytes(wb), 1)
    assert str(e.value) == MSG_TOO_MANY_CELLS


def test_cell_limit_boundary_ok():
    wb = Workbook()
    ws = wb.active
    for i in range(1, MAX_XLSX_CELLS + 1):
        ws.cell(row=i, column=1, value="x")
    assert len(parse_xlsx(to_bytes(wb), 1).nodes) == MAX_XLSX_CELLS


def test_empty_workbook():
    wb = Workbook()
    wb.active["A1"] = "   "
    with pytest.raises(ParseError) as e:
        parse_xlsx(to_bytes(wb), 1)
    assert str(e.value) == MSG_EMPTY


@pytest.mark.parametrize("data", [b"", b"not a zip", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"secret-text" * 50])
def test_corrupt_or_encrypted(data):
    with pytest.raises(ParseError) as e:
        parse_xlsx(data, 1)
    assert str(e.value) == MSG_CORRUPT
    assert "secret" not in str(e.value)


def test_truncated_zip_message_has_no_content():
    data = (FIXTURES / "quotation.xlsx").read_bytes()
    with pytest.raises(ParseError) as e:
        parse_xlsx(data[: len(data) // 2], 1)
    assert str(e.value) == MSG_CORRUPT
