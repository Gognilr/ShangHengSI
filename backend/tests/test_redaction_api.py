"""M2-13、M4-05：真实脱敏确认接口、重算、疑似遗漏决定与手动补充、状态限制及 A-06 归属隔离。"""
import io
import logging
import uuid
from collections import Counter

import docx
import pytest
from sqlalchemy import select

from app.db.models import DocNode, LlmCall, Party, RedactionItem, RedactionTerm, Task
from app.main import API_PREFIX
from app.redaction import engine
from app.security.crypto import decrypt_text
from tests.review_support import REAL_BUYER_NAME, REAL_OUR_NAME, env, llm, ready_task  # noqa: F401

pytestmark = pytest.mark.usefixtures("env")


@pytest.fixture
def alice(make_user, login):
    user = make_user("alice")
    client, csrf = login("alice")
    return user, client, {"X-CSRF-Token": csrf}


def url(task, suffix):
    return f"{API_PREFIX}/tasks/{task.id}/{suffix}"


@pytest.mark.parametrize("mode", ["manual", "suspect"])
def test_v2c06_remember_custom_immediate_and_deduplicated(alice,db_sessionmaker,ready_task,mode):
    from app.db.models import RedactionSuspect
    from app.security.crypto import encrypt_text
    user,c,h = alice
    task = ready_task(user)
    duplicate_task = ready_task(user, user_note="示例简称用于另一任务")
    with db_sessionmaker() as db:
        node = db.scalar(select(DocNode).where(DocNode.task_id==task.id,DocNode.part=="body"))
        node.text_enc = encrypt_text("示例简称用于本次报价。")
        loc = node.loc
        db.commit()
    if mode == "manual":
        response = c.post(url(task,"redaction/manual"),json={"loc":loc,"char_start":0,"char_end":4,"category":"CUSTOM","remember":True},headers=h)
    else:
        with db_sessionmaker() as db:
            suspect = RedactionSuspect(task_id=task.id,loc=loc,char_start=0,char_end=4,reason="capitalized")
            db.add(suspect)
            db.commit()
            sid = suspect.id
        response = c.post(url(task,f"redaction/suspects/{sid}"),json={"decision":"replace","category":"CUSTOM","remember":True},headers=h)
    assert response.status_code == 200,response.text
    duplicate = c.post(url(duplicate_task,"redaction/manual"),
                       json={"loc":"NOTE","char_start":0,"char_end":4,"category":"CUSTOM","remember":True}, headers=h)
    assert duplicate.status_code == 200, duplicate.text
    with db_sessionmaker() as db:
        terms = list(db.scalars(select(RedactionTerm).where(RedactionTerm.kind=="custom")))
        assert len(terms)==1 and decrypt_text(terms[0].term_enc)=="示例简称"
        assert terms[0].category=="CUSTOM" and terms[0].enabled and terms[0].created_by==user.id
    next_task = ready_task(user,user_note="示例简称用于补充要求")
    with db_sessionmaker() as db:
        assert "示例简称" not in (db.get(Task,next_task.id).user_note_redacted or "")
        assert len(list(db.scalars(select(RedactionTerm).where(RedactionTerm.kind=="custom")))) == 1


@pytest.mark.parametrize("mode", ["manual", "suspect"])
def test_v2c06_remember_prompt_conflict_is_atomic(alice,db_sessionmaker,ready_task,mode):
    from app.db.models import RedactionSuspect
    from app.security.crypto import encrypt_text
    user,c,h = alice
    task = ready_task(user)
    with db_sessionmaker() as db:
        node = db.scalar(select(DocNode).where(DocNode.task_id==task.id,DocNode.part=="body"))
        node.text_enc = encrypt_text("待审核文件")
        loc = node.loc
        db.commit()
        before = [(i.loc,i.char_start,i.char_end,i.placeholder) for i in db.scalars(select(RedactionItem).where(RedactionItem.task_id==task.id))]
    if mode == "manual":
        response = c.post(url(task,"redaction/manual"),json={"loc":loc,"char_start":0,"char_end":5,"category":"CUSTOM","remember":True},headers=h)
    else:
        with db_sessionmaker() as db:
            suspect = RedactionSuspect(task_id=task.id,loc=loc,char_start=0,char_end=5,reason="capitalized")
            db.add(suspect)
            db.commit()
            sid = suspect.id
        response = c.post(url(task,f"redaction/suspects/{sid}"),json={"decision":"replace","category":"CUSTOM","remember":True},headers=h)
    assert response.status_code==422 and response.json()["code"]=="TERM_IN_PROMPT"
    with db_sessionmaker() as db:
        assert list(db.scalars(select(RedactionTerm).where(RedactionTerm.kind=="custom")))==[]
        assert [(i.loc,i.char_start,i.char_end,i.placeholder) for i in db.scalars(select(RedactionItem).where(RedactionItem.task_id==task.id))]==before
        if mode == "suspect":
            assert db.get(RedactionSuspect, sid).decision == "pending"


def test_v2c06_remember_does_not_enable_disabled_term(alice,db_sessionmaker,ready_task):
    from app.security.crypto import encrypt_text
    user,c,h = alice
    task = ready_task(user, user_note="示例简称用于本次补充")
    with db_sessionmaker() as db:
        term = RedactionTerm(kind="custom",category="CUSTOM",term_hash=engine.term_hash("示例简称"),
                             term_enc=encrypt_text("示例简称"),enabled=False,created_by=None)
        db.add(term)
        db.commit()
        key = term.id
    response = c.post(url(task,"redaction/manual"),
                      json={"loc":"NOTE","char_start":0,"char_end":4,"category":"CUSTOM","remember":True},headers=h)
    assert response.status_code == 200, response.text
    with db_sessionmaker() as db:
        assert not db.get(RedactionTerm,key).enabled
        assert db.get(RedactionTerm,key).created_by is None
        assert len(list(db.scalars(select(RedactionTerm).where(RedactionTerm.kind=="custom")))) == 1


def parties(client, task):
    r = client.get(url(task, "parties"))
    assert r.status_code == 200, r.text
    return r.json()["items"]


def payload(items, our_id=None):
    return {"items": [
        {"id": p["id"], "is_our_party": p["id"] == our_id if our_id else p["is_our_party"],
         "country": p["country"], "org_type": p["org_type"]} for p in items
    ]}


def file_id(client, task):
    return client.get(f"{API_PREFIX}/tasks/{task.id}").json()["files"][0]["id"]


def test_get_parties_summary_outbound(db_sessionmaker, alice, ready_task):
    user, c, _ = alice
    task = ready_task(user)
    items = parties(c, task)
    assert {p["name"] for p in items} == {REAL_OUR_NAME, REAL_BUYER_NAME}
    assert sum(p["is_our_party"] for p in items) == 1
    assert all(set(p) == {"id", "code", "name", "role_label", "country", "org_type", "is_our_party"}
               for p in items)
    r = c.get(url(task, "redaction"))
    assert r.status_code == 200
    data = r.json()
    assert data["suspects"] == []
    assert data["summary"] == dict(Counter(i["category"] for i in data["items"]))
    assert data["summary"]["ORG"] >= 2 and data["summary"]["AMOUNT"] >= 1
    assert all(set(i) == {"loc", "char_start", "char_end", "category", "placeholder", "can_revoke"}
               and i["char_start"] < i["char_end"] for i in data["items"])
    assert all(isinstance(i["can_revoke"], bool) for i in data["items"])
    r = c.get(url(task, "redaction/outbound"))
    assert r.status_code == 200
    text = r.json()["text"]
    assert "【主体说明】" in text and "[F1-" in text
    assert REAL_OUR_NAME not in text and REAL_BUYER_NAME not in text
    assert "我方" in text and "客户_A" in text
    with db_sessionmaker() as db:
        assert text == engine.build_outbound_text(db, db.get(Task, task.id))
        assert list(db.scalars(select(LlmCall))) == []


@pytest.mark.parametrize("reverse", [False, True])
def test_put_rebuilds_codes_map_nodes_and_persists(db_sessionmaker, alice, ready_task, reverse):
    user, c, h = alice
    task = ready_task(user)
    before = parties(c, task)
    buyer = next(p for p in before if p["name"] == REAL_BUYER_NAME)
    body = payload(before, buyer["id"])
    for p in body["items"]:
        if p["id"] == buyer["id"]:
            p.update(country="土耳其", org_type="虚构自定义类型")
        else:
            p.update(country="德国", org_type="货代")
    if reverse:
        body["items"].reverse()
    r = c.put(url(task, "parties"), json=body, headers=h)
    assert r.status_code == 200, r.text
    after = r.json()["items"]
    assert {p["id"] for p in after} == {p["id"] for p in before}
    assert sum(p["is_our_party"] for p in after) == 1
    changed = next(p for p in after if p["id"] == buyer["id"])
    assert (changed["code"], changed["country"], changed["org_type"]) == ("我方", "土耳其", "虚构自定义类型")
    assert parties(c, task) == after
    text = c.get(url(task, "redaction/outbound")).json()["text"]
    assert "我方=买方" in text and "德国，货代" in text
    assert REAL_OUR_NAME not in text and REAL_BUYER_NAME not in text
    with db_sessionmaker() as db:
        assert engine.load_redaction_map(db, task.id)["我方"] == REAL_BUYER_NAME
        assert all(p.confirmed for p in db.scalars(select(Party).where(Party.task_id == task.id)))
        buyer_nodes = [n for n in db.scalars(select(DocNode).where(DocNode.task_id == task.id))
                       if n.text_enc and REAL_BUYER_NAME in decrypt_text(n.text_enc)]
        assert buyer_nodes and all("我方" in n.redacted_text for n in buyer_nodes)
    # 同样的完整列表可重复提交。
    assert c.put(url(task, "parties"), json=body, headers=h).status_code == 200


@pytest.mark.parametrize("mode", ["none", "both", "empty"])
def test_exactly_one_our_party(alice, ready_task, mode):
    user, c, h = alice
    task = ready_task(user)
    before = parties(c, task)
    body = payload(before)
    if mode == "empty":
        body["items"] = []
        expected = "OUR_PARTY_REQUIRED"
    else:
        for p in body["items"]:
            p["is_our_party"] = mode == "both"
        expected = "OUR_PARTY_REQUIRED"
    r = c.put(url(task, "parties"), json=body, headers=h)
    assert (r.status_code, r.json()["code"]) == (422, expected)
    assert parties(c, task) == before


@pytest.mark.parametrize("mode", ["duplicate", "missing", "foreign", "unknown", "country", "type", "blank", "bool"])
def test_invalid_party_input_is_atomic(alice, ready_task, mode):
    user, c, h = alice
    task = ready_task(user)
    before = parties(c, task)
    body = payload(before)
    if mode == "duplicate":
        body["items"].append(body["items"][0].copy())
    elif mode == "missing":
        body["items"].pop()
    elif mode == "foreign":
        other = ready_task(user)
        body["items"][0]["id"] = parties(c, other)[0]["id"]
    elif mode == "unknown":
        body["items"][0]["id"] = str(uuid.uuid4())
    elif mode == "country":
        body["items"][0]["country"] = "虚" * 51
    elif mode == "type":
        body["items"][0]["org_type"] = "虚" * 21
    elif mode == "blank":
        body["items"][0]["country"] = "  "
    else:
        body["items"][0]["is_our_party"] = "false"
    r = c.put(url(task, "parties"), json=body, headers=h)
    assert (r.status_code, r.json()["code"]) == (422, "VALIDATION_ERROR")
    assert parties(c, task) == before


@pytest.mark.parametrize("state", ["parsing", "parse_failed", "reviewing", "reviewed", "completed_local"])
def test_mutations_only_awaiting_confirm(db_sessionmaker, alice, ready_task, state):
    user, c, h = alice
    task = ready_task(user)
    before = parties(c, task)
    fid = file_id(c, task)
    with db_sessionmaker() as db:
        db.get(Task, task.id).status = state
        db.commit()
    for r in [c.put(url(task, "parties"), json=payload(before), headers=h),
              c.patch(url(task, f"files/{fid}"), json={"doc_type": "letter"}, headers=h)]:
        assert (r.status_code, r.json()["code"]) == (409, "TASK_STATE_CONFLICT")
    assert parties(c, task) == before
    assert c.get(f"{API_PREFIX}/tasks/{task.id}").json()["files"][0]["doc_type_source"] == "auto"


@pytest.mark.parametrize("doc_type", ["contract", "quotation", "order", "inquiry", "letter"])
def test_patch_file_type(alice, ready_task, doc_type):
    user, c, h = alice
    task = ready_task(user)
    fid = file_id(c, task)
    before = c.get(f"{API_PREFIX}/tasks/{task.id}").json()["files"][0]
    r = c.patch(url(task, f"files/{fid}"), json={"doc_type": doc_type}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json() == dict(before, doc_type=doc_type, doc_type_source="user")
    assert c.get(f"{API_PREFIX}/tasks/{task.id}").json()["files"][0] == r.json()


@pytest.mark.parametrize("body", [{}, {"doc_type": None}, {"doc_type": "other"}])
def test_patch_invalid_type(alice, ready_task, body):
    user, c, h = alice
    task = ready_task(user)
    r = c.patch(url(task, f"files/{file_id(c, task)}"), json=body, headers=h)
    assert (r.status_code, r.json()["code"]) == (422, "VALIDATION_ERROR")


@pytest.mark.parametrize("role", ["user", "admin"])
def test_a06_other_user_all_real_endpoints(make_user, login, alice, ready_task, role):
    owner, c, _ = alice
    task = ready_task(owner)
    body = payload(parties(c, task))
    fid = file_id(c, task)
    make_user("intruder", role=role)
    stranger, csrf = login("intruder")
    h = {"X-CSRF-Token": csrf}
    for tid in [str(task.id), str(uuid.uuid4()), "invalid-id"]:
        base = f"{API_PREFIX}/tasks/{tid}"
        responses = [stranger.get(base + "/parties"), stranger.get(base + "/redaction"),
                     stranger.get(base + "/redaction/outbound"),
                     stranger.put(base + "/parties", json=body, headers=h),
                     stranger.patch(base + f"/files/{fid}", json={"doc_type": "letter"}, headers=h)]
        for r in responses:
            assert (r.status_code, r.json()["code"]) == (404, "TASK_NOT_FOUND")


def test_cross_task_file_and_csrf(alice, ready_task):
    user, c, h = alice
    task = ready_task(user)
    other = ready_task(user)
    for fid in [file_id(c, other), str(uuid.uuid4()), "invalid-id"]:
        r = c.patch(url(task, f"files/{fid}"), json={"doc_type": "letter"}, headers=h)
        assert (r.status_code, r.json()["code"]) == (404, "FILE_NOT_FOUND")
    r = c.put(url(task, "parties"), json=payload(parties(c, task)))
    assert (r.status_code, r.json()["code"]) == (403, "CSRF_FAILED")
    r = c.patch(url(task, f"files/{file_id(c, task)}"), json={"doc_type": "letter"})
    assert (r.status_code, r.json()["code"]) == (403, "CSRF_FAILED")


def test_errors_and_logs_do_not_echo_input(alice, ready_task, caplog):
    user, c, h = alice
    task = ready_task(user)
    body = payload(parties(c, task))
    secret = "虚构敏感名称金额888888密钥secret"
    body["items"][0]["country"] = secret * 4
    r = c.put(url(task, "parties"), json=body, headers=h)
    assert r.status_code == 422
    assert secret not in r.text + caplog.text
    assert REAL_OUR_NAME not in caplog.text and REAL_BUYER_NAME not in caplog.text


# ---------------------------------------------------------------- M4-05：疑似遗漏决定、手动补充（设计 M4.md 第五节）


def docx_of(*paragraphs: str) -> bytes:
    doc = docx.Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


SUSPECT_DOC = (
    "销售确认书",
    f"卖方：{REAL_OUR_NAME}",
    f"买方：{REAL_BUYER_NAME}",
    "货物由 Northwind Fictive Holdings 承运。",
    "Northwind Fictive Holdings 另行通知装船日期。",
    "收款 6222 0000 1111 2222 333，请核对。",
)


def suspect_task(ready_task, user, *paragraphs):
    return ready_task(user, [("确认书.docx", docx_of(*(paragraphs or SUSPECT_DOC)))])


def redaction(client, task):
    r = client.get(url(task, "redaction"))
    assert r.status_code == 200, r.text
    return r.json()


def node_loc(client, task, needle):
    """正文中含 needle 的节点 loc 与 needle 在节点原文中的位置。"""
    r = client.get(f"{API_PREFIX}/tasks/{task.id}/files/{file_id(client, task)}/nodes")
    assert r.status_code == 200, r.text
    node = next(n for n in r.json()["items"] if n["part"] == "body" and needle in (n["text"] or ""))
    start = node["text"].index(needle)
    return node["loc"], start, start + len(needle)


def test_get_redaction_lists_suspects_with_context(alice, ready_task):
    user, c, _ = alice
    task = suspect_task(ready_task, user)
    data = redaction(c, task)
    views = data["suspects"]
    assert [(v["reason"], v["text"], v["decision"]) for v in views] == [
        ("capitalized", "Northwind Fictive Holdings", "pending"),
        ("long_digits", "6222 0000 1111 2222 333", "pending")]
    assert [v["occurrences"] for v in views] == [2, 1]
    assert all(set(v) == {"id", "loc", "reason", "decision", "category", "text", "before", "after", "occurrences"} for v in views)
    assert (views[0]["before"], views[0]["after"]) == ("货物由 ", " 承运。")
    assert views[0]["loc"] == node_loc(c, task, "货物由 Northwind Fictive Holdings")[0]
    account = next(v for v in views if v["reason"] == "long_digits" and v["text"] == "6222 0000 1111 2222 333")
    assert account["before"] == "收款 " and account["category"] is None


def test_f03_batch_ignore_is_owned_and_atomic(alice, ready_task, make_user, db_sessionmaker):
    from app.db.models import RedactionSuspect

    user, c, h = alice
    task = suspect_task(ready_task, user)
    other = suspect_task(ready_task, make_user("batch_other"))
    before = redaction(c, task)
    ids = [s["id"] for s in before["suspects"]]
    assert len(ids) >= 2
    with db_sessionmaker() as db:
        foreign = db.scalar(select(RedactionSuspect.id).where(RedactionSuspect.task_id == other.id))
    endpoint = url(task, "redaction/suspects/batch")
    r = c.post(endpoint, json={"ids": [ids[0], foreign], "decision": "ignore"}, headers=h)
    assert r.status_code == 404, r.text
    assert redaction(c, task) == before
    r = c.post(endpoint, json={"ids": ids, "decision": "ignore"}, headers=h)
    assert r.status_code == 200, r.text
    assert all(s["decision"] == "ignore" for s in r.json()["suspects"])
    assert r.json() == redaction(c, task)
    assert r.json()["summary"] == before["summary"]


def test_ignore_then_replace_and_send(db_sessionmaker, alice, ready_task, llm):
    user, c, h = alice
    task = suspect_task(ready_task, user)
    views = redaction(c, task)["suspects"]
    # 确认发送被未处理的疑似遗漏拦住
    assert c.post(url(task, "runs"), headers=h).json()["code"] == "SUSPECTS_PENDING"
    # 忽略：同规范名称的所有原因一并忽略，视图仍按组返回（B-12）。
    r = c.post(url(task, f"redaction/suspects/{views[0]['id']}"), json={"decision": "ignore"}, headers=h)
    assert r.status_code == 200, r.text
    after = r.json()["suspects"]
    assert [(v["text"], v["decision"]) for v in after] == [
        ("6222 0000 1111 2222 333", "pending"), ("Northwind Fictive Holdings", "ignore")]
    assert [v["occurrences"] for v in after] == [1, 2]
    assert r.json() == redaction(c, task)
    # 已忽略的可以改为替换：全文替换，作为新的一方出现在签约方列表中
    r = c.post(url(task, f"redaction/suspects/{after[1]['id']}"), json={"decision": "replace", "category": "ORG"},
               headers=h)
    assert r.status_code == 200, r.text
    assert [v["reason"] for v in r.json()["suspects"]] == ["long_digits"]
    assert r.json()["summary"]["ORG"] >= 3
    names = {p["name"]: p["code"] for p in parties(c, task)}
    assert names["Northwind Fictive Holdings"] == "公司_B"
    sid = r.json()["suspects"][0]["id"]
    r = c.post(url(task, f"redaction/suspects/{sid}"), json={"decision": "replace", "category": "CUSTOM"}, headers=h)
    assert r.status_code == 200 and r.json()["suspects"] == []
    out = c.get(url(task, "redaction/outbound")).json()["text"]
    assert "Northwind" not in out and "6222" not in out and "[敏感词_1]" in out
    with db_sessionmaker() as db:
        assert list(db.scalars(select(LlmCall))) == []  # 第二步的操作不调用模型
    llm("normal")
    r = c.post(url(task, "runs"), headers=h)
    assert r.status_code == 201, r.text


def test_suspect_errors(db_sessionmaker, alice, ready_task):
    user, c, h = alice
    task = suspect_task(ready_task, user)
    other = suspect_task(ready_task, user)
    mine = redaction(c, task)["suspects"][0]["id"]
    theirs = redaction(c, other)["suspects"][0]["id"]
    for sid in (theirs, 10 ** 12, "abc", "-1", str(2 ** 70)):
        r = c.post(url(task, f"redaction/suspects/{sid}"), json={"decision": "ignore"}, headers=h)
        assert (r.status_code, r.json()["code"]) == (404, "SUSPECT_NOT_FOUND"), sid
    for body in ({"decision": "replace"}, {"decision": "replace", "category": None},
                 {"decision": "replace", "category": "OUR"}, {"decision": "delete"}, {}):
        r = c.post(url(task, f"redaction/suspects/{mine}"), json=body, headers=h)
        assert (r.status_code, r.json()["code"]) == (422, "VALIDATION_ERROR"), body
    assert {v["decision"] for v in redaction(c, task)["suspects"]} == {"pending"}
    assert {v["decision"] for v in redaction(c, other)["suspects"]} == {"pending"}


def test_manual_errors(alice, ready_task):
    user, c, h = alice
    task = ready_task(user)
    before = redaction(c, task)
    loc, start, end = node_loc(c, task, "煅后石油焦")
    our_loc, our_start, our_end = node_loc(c, task, REAL_OUR_NAME)

    def post(**kw):
        body = {"loc": loc, "char_start": start, "char_end": end, "category": "CUSTOM", **kw}
        return c.post(url(task, "redaction/manual"), json=body, headers=h)

    cases = [
        ({"loc": "F9-P1"}, 422, "INVALID_SELECTION"),
        ({"char_start": -1}, 422, "INVALID_SELECTION"),
        ({"char_end": 10 ** 6}, 422, "INVALID_SELECTION"),
        ({"char_start": end, "char_end": start}, 422, "INVALID_SELECTION"),
        ({"loc": our_loc, "char_start": our_start, "char_end": our_end}, 409, "ALREADY_REDACTED"),
        ({"loc": our_loc, "char_start": max(0, our_start - 2), "char_end": our_start + 2}, 409, "ALREADY_REDACTED"),
        ({"char_end": start + 1}, 422, "SELECTION_LENGTH"),
        ({"category": "OUR"}, 422, "VALIDATION_ERROR"),
        ({"category": None}, 422, "VALIDATION_ERROR"),
        ({"char_start": "1"}, 422, "VALIDATION_ERROR"),
        ({"loc": ""}, 422, "VALIDATION_ERROR"),
    ]
    for kw, status, code in cases:
        r = post(**kw)
        assert (r.status_code, r.json()["code"]) == (status, code), kw
    assert redaction(c, task) == before
    r = post()
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["CUSTOM"] >= 1
    assert "煅后石油焦" not in c.get(url(task, "redaction/outbound")).json()["text"]


TERM_IN_PROMPT = (422, "TERM_IN_PROMPT")


def test_manual_rejects_text_in_fixed_prompt(db_sessionmaker, alice, ready_task):
    """Q-77 ③：手动补充的文字出现在审核提示中（公司说明、内置要求）时拒绝，具体名称不受影响。"""
    from app.admin import settings_store

    with db_sessionmaker() as db:
        settings_store.set_setting(db, "review.company_profile", "本公司主营虚构钛白粉出口", None)
        db.commit()
    user, c, h = alice
    task = suspect_task(ready_task, user, "销售确认书", f"卖方：{REAL_OUR_NAME}", f"买方：{REAL_BUYER_NAME}",
                        "货物：虚构钛白粉 20 吨，经 Lumen Fictive Forwarding 代为装运。", "违约金及交货期见附件。")
    before = redaction(c, task)
    for needle in ("虚构钛白粉", "违约金", "交货期"):
        loc, start, end = node_loc(c, task, needle)
        r = c.post(url(task, "redaction/manual"), json={"loc": loc, "char_start": start, "char_end": end,
                                                        "category": "CUSTOM"}, headers=h)
        assert (r.status_code, r.json()["code"]) == TERM_IN_PROMPT, needle
        assert r.json()["message"] == "所选文字在审核提示中也会出现，设为敏感信息后将无法发送，请只选择具体的名称或编号"
        assert needle not in r.text
    assert redaction(c, task) == before
    loc, start, end = node_loc(c, task, "Lumen Fictive Forwarding")
    r = c.post(url(task, "redaction/manual"), json={"loc": loc, "char_start": start, "char_end": end,
                                                    "category": "ORG"}, headers=h)
    assert r.status_code == 200, r.text
    assert "Lumen" not in c.get(url(task, "redaction/outbound")).json()["text"]


def test_suspect_replace_rejects_text_in_fixed_prompt(db_sessionmaker, alice, ready_task):
    """Q-77 ③：疑似遗漏选“替换”时同样拦截（文字出现在启用的公司补充要求中）；忽略不受影响，停用的要求不算。"""
    from app.db.models import CompanyRequirement

    with db_sessionmaker() as db:
        db.add_all([CompanyRequirement(code="CO-1", text="与 Northwind Fictive Holdings 的往来须经法务确认"),
                    CompanyRequirement(code="CO-2", text="6222 0000 1111 2222 333", enabled=False)])
        db.commit()
    user, c, h = alice
    task = suspect_task(ready_task, user)
    before = redaction(c, task)
    views = before["suspects"]
    r = c.post(url(task, f"redaction/suspects/{views[0]['id']}"), json={"decision": "replace", "category": "ORG"},
               headers=h)
    assert (r.status_code, r.json()["code"]) == TERM_IN_PROMPT
    assert "Northwind" not in r.text
    assert redaction(c, task) == before  # 决定未写入
    account = next(v for v in views if v["reason"] == "long_digits" and v["text"] == "6222 0000 1111 2222 333")
    r = c.post(url(task, f"redaction/suspects/{account['id']}"), json={"decision": "replace", "category": "CUSTOM"},
               headers=h)
    assert r.status_code == 200, r.text
    sid = next(v["id"] for v in r.json()["suspects"] if v["text"] == "Northwind Fictive Holdings")
    r = c.post(url(task, f"redaction/suspects/{sid}"), json={"decision": "ignore"}, headers=h)
    assert r.status_code == 200, r.text
    assert [(v["text"], v["decision"], v["occurrences"]) for v in r.json()["suspects"]] == [
        ("Northwind Fictive Holdings", "ignore", 2)]


def test_manual_selection_length_upper_bound(alice, ready_task):
    user, c, h = alice
    task = suspect_task(ready_task, user, "x" * 150, f"卖方：{REAL_OUR_NAME}")
    loc, start, _ = node_loc(c, task, "x" * 150)
    r = c.post(url(task, "redaction/manual"), json={"loc": loc, "char_start": 0, "char_end": 101, "category": "CUSTOM"},
               headers=h)
    assert (r.status_code, r.json()["code"]) == (422, "SELECTION_LENGTH")


def test_q54_manual_party_then_our_party_and_send(db_sessionmaker, alice, ready_task, llm):
    """Q-54：没有识别出签约方的文件，手动补充公司名后设为我方并确认发送。"""
    user, c, h = alice
    task = suspect_task(ready_task, user, "合作备忘录", "本备忘录由远景示例科技与对方共同确认。")
    assert parties(c, task) == []
    assert c.post(url(task, "runs"), headers=h).json()["code"] in ("SUSPECTS_PENDING", "OUR_PARTY_REQUIRED")
    for v in redaction(c, task)["suspects"]:
        c.post(url(task, f"redaction/suspects/{v['id']}"), json={"decision": "ignore"}, headers=h)
    loc, start, end = node_loc(c, task, "远景示例科技")
    r = c.post(url(task, "redaction/manual"), json={"loc": loc, "char_start": start, "char_end": end, "category": "ORG"},
               headers=h)
    assert r.status_code == 200, r.text
    items = parties(c, task)
    assert [(p["name"], p["code"], p["role_label"], p["is_our_party"]) for p in items] == [
        ("远景示例科技", "公司_A", "其他", False)]
    r = c.put(url(task, "parties"), json=payload(items, items[0]["id"]), headers=h)
    assert r.status_code == 200, r.text
    assert "本备忘录由我方与对方共同确认。" in c.get(url(task, "redaction/outbound")).json()["text"]
    llm("normal")
    assert c.post(url(task, "runs"), headers=h).status_code == 201


def test_t08_learned_name_replaced_in_second_task(db_sessionmaker, alice, ready_task, llm):
    """T-08 的接口部分：手动补充的公司名在确认发送后记住，第二个任务自动替换。"""
    user, c, h = alice
    first = suspect_task(ready_task, user, f"卖方：{REAL_OUR_NAME}", f"买方：{REAL_BUYER_NAME}", "请转交青松示例贸易处理。")
    loc, start, end = node_loc(c, first, "青松示例贸易")
    assert c.post(url(first, "redaction/manual"),
                  json={"loc": loc, "char_start": start, "char_end": end, "category": "ORG"}, headers=h).status_code == 200
    for v in redaction(c, first)["suspects"]:
        c.post(url(first, f"redaction/suspects/{v['id']}"), json={"decision": "ignore"}, headers=h)
    llm("normal")
    assert c.post(url(first, "runs"), headers=h).status_code == 201
    with db_sessionmaker() as db:
        learned = {decrypt_text(t.term_enc): t.is_our_party
                   for t in db.scalars(select(RedactionTerm).where(RedactionTerm.kind == "learned"))}
    assert learned == {REAL_OUR_NAME: True, REAL_BUYER_NAME: False, "青松示例贸易": False}
    second = suspect_task(ready_task, user, "通知", "青松示例贸易将于下月到访。")
    out = c.get(url(second, "redaction/outbound")).json()["text"]
    assert "青松示例贸易" not in out and "[F1-P2] 公司_A将于下月到访。" in out
    with db_sessionmaker() as db:
        items = db.scalars(select(RedactionItem).where(RedactionItem.task_id == second.id)).all()
        assert {i.detected_by for i in items} == {"term"}
        hit = db.scalar(select(RedactionTerm).where(RedactionTerm.term_hash == engine.term_hash("青松示例贸易")))
        assert hit.hit_count == 1 and hit.last_hit_at is not None


@pytest.mark.parametrize("state", ["parsing", "parse_failed", "reviewing", "reviewed", "completed_local"])
def test_new_mutations_only_awaiting_confirm(db_sessionmaker, alice, ready_task, state):
    user, c, h = alice
    task = suspect_task(ready_task, user)
    before = redaction(c, task)
    sid = before["suspects"][0]["id"]
    loc, start, end = node_loc(c, task, "销售确认书")
    with db_sessionmaker() as db:
        db.get(Task, task.id).status = state
        db.commit()
    for r in [c.post(url(task, f"redaction/suspects/{sid}"), json={"decision": "ignore"}, headers=h),
              c.post(url(task, "redaction/manual"),
                     json={"loc": loc, "char_start": start, "char_end": end, "category": "CUSTOM"}, headers=h)]:
        assert (r.status_code, r.json()["code"]) == (409, "TASK_STATE_CONFLICT")
    assert redaction(c, task) == before


@pytest.mark.parametrize("role", ["user", "admin"])
def test_a06_new_endpoints(make_user, login, alice, ready_task, role):
    owner, c, h_owner = alice
    task = suspect_task(ready_task, owner)
    sid = redaction(c, task)["suspects"][0]["id"]
    make_user("intruder", role=role)
    stranger, csrf = login("intruder")
    h = {"X-CSRF-Token": csrf}
    for tid in [str(task.id), str(uuid.uuid4()), "invalid-id"]:
        base = f"{API_PREFIX}/tasks/{tid}"
        for r in [stranger.post(base + f"/redaction/suspects/{sid}", json={"decision": "ignore"}, headers=h),
                  stranger.post(base + "/redaction/manual",
                                json={"loc": "F1-P1", "char_start": 0, "char_end": 3, "category": "CUSTOM"}, headers=h)]:
            assert (r.status_code, r.json()["code"]) == (404, "TASK_NOT_FOUND")
    assert {v["decision"] for v in redaction(c, task)["suspects"]} == {"pending"}
    # 不带 CSRF 令牌
    r = c.post(url(task, f"redaction/suspects/{sid}"), json={"decision": "ignore"})
    assert (r.status_code, r.json()["code"]) == (403, "CSRF_FAILED")
    r = c.post(url(task, "redaction/manual"), json={"loc": "F1-P1", "char_start": 0, "char_end": 3, "category": "CUSTOM"})
    assert (r.status_code, r.json()["code"]) == (403, "CSRF_FAILED")


def test_new_endpoints_do_not_log_text(alice, ready_task, caplog):
    user, c, h = alice
    task = suspect_task(ready_task, user)
    caplog.set_level(logging.DEBUG)
    views = redaction(c, task)["suspects"]
    c.post(url(task, f"redaction/suspects/{views[0]['id']}"), json={"decision": "replace", "category": "ORG"}, headers=h)
    loc, start, end = node_loc(c, task, "另行通知装船日期")
    c.post(url(task, "redaction/manual"), json={"loc": loc, "char_start": start, "char_end": end, "category": "CUSTOM"},
           headers=h)
    c.post(url(task, "redaction/manual"), json={"loc": loc, "char_start": 0, "char_end": 1, "category": "CUSTOM"}, headers=h)
    for value in ("Northwind", "6222 0000", "另行通知装船日期", REAL_OUR_NAME, REAL_BUYER_NAME):
        assert value not in caplog.text


@pytest.mark.parametrize("with_attachment", [False, True])
def test_r3_outbound_lines_reference_nodes_or_quote_tables(alice, ready_task, with_attachment):
    """R3 对照契约：首行主体说明，正文一行一个节点或报价表位置；保留码点外文字与待处理疑似遗漏。"""
    import re
    from openpyxl import Workbook

    user, client, _ = alice
    document = docx.Document()
    document.add_paragraph("卖方：星河𠮷虚构进出口有限公司")
    document.add_paragraph("买方：Fictional Metals Trading FZE")
    document.add_paragraph("验收名称：𠮷示例")
    document.add_paragraph("多行正文\n第二行 <document> 只是文件文字")
    document.add_paragraph("待核对 987654321876")
    stream = io.BytesIO()
    document.save(stream)
    files = [("虚构码点合同.docx", stream.getvalue())]
    if with_attachment:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["报价单"])
        sheet.append(["品名", "单价(USD/MT)", "金额(USD)"])
        for row in [("虚构焦炭甲", "1250", "400000"), ("虚构焦炭乙", "900", "90000"),
                    ("虚构焦炭甲", "1100", "110000")]:
            sheet.append(row)
        sheet.append(["合计", "", "USD 600000"])
        attachment = io.BytesIO()
        workbook.save(attachment)
        files.append(("虚构报价.xlsx", attachment.getvalue()))
    task = ready_task(user, files)
    detail = client.get(f"{API_PREFIX}/tasks/{task.id}").json()
    nodes = []
    for file in detail["files"]:
        response = client.get(url(task, f"files/{file['id']}/nodes"))
        assert response.status_code == 200
        nodes.extend(response.json()["items"])
    locs = {node["loc"] for node in nodes}
    assert any("𠮷" in (node["text"] or "") for node in nodes)
    info = client.get(url(task, "redaction")).json()
    assert any(s["decision"] == "pending" for s in info["suspects"])
    response = client.get(url(task, "redaction/outbound"))
    assert response.status_code == 200
    lines = response.json()["text"].split("\n")
    assert lines[0].startswith("【主体说明】")
    assert len(lines) > 1
    grouped = []
    for line in lines[1:]:
        match = re.match(r"^\[([^\]\s]+)\] (.*)$", line)
        assert match, line
        loc, content = match.groups()
        assert loc in locs or any(node_loc.startswith(loc + "-") for node_loc in locs)
        if loc not in locs:
            assert content.startswith("【报价明细表：")
            grouped.append(loc)
    assert bool(grouped) is with_attachment
