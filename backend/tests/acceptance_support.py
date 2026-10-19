"""M4 验收（M4-07）的公共部分：经真实接口把一份样本从上传走到审核完成。

流程与浏览器中一致：POST /tasks 上传 → （代替 worker）解析 → GET /redaction → 逐条处理疑似遗漏 →
没有我方时补充并设定我方 → POST /runs 确认发送 → （代替 worker）用模拟模型执行审核 → GET 轮次详情。
模型调用一律经 tests.review_support 的 llm 夹具（httpx.MockTransport），不会联网。
"""
import io
import json
import uuid
from dataclasses import dataclass, field
from pathlib import Path

import docx
import openpyxl
import yaml
from sqlalchemy import select, update

from app.db.models import Job, LlmCall, RedactionTerm, Task
from app.main import API_PREFIX
from app.redaction.engine import load_real_values
from app.review import job
from app.security.crypto import decrypt_text
from app.tasks import pipeline

FIXTURES = Path(__file__).resolve().parent / "fixtures"
GOLDEN = FIXTURES / "redaction_golden"
# 解析器能处理的 fixtures 目录（encrypted.pdf 预期解析失败，单独排除）
FIXTURE_DIRS = ("docx", "pdf", "xlsx", "images")
PARSE_FAIL_EXPECTED = {"encrypted.pdf"}
# fixtures 标注里与 SPEC 9.2 不一致的 must_replace：REFNO 只含合同号、订单号、发票号，
# 图片报价函中“编号：QT-2026-0508”是报价单号（PDF 样本把同一编号标为 must_keep），记入 open_questions Q-77
NOT_IN_SPEC = {"QT-2026-0508"}


@dataclass
class Sample:
    name: str  # 上传时的文件名
    data: bytes
    stance: str = "sales"
    must_replace: list[str] = field(default_factory=list)
    must_suspect: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- 样本


def _docx_bytes(lines: list[str]) -> bytes:
    d = docx.Document()
    for line in lines:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes(cells: list[tuple[str, str]]) -> bytes:
    """cells: [("F1-S1-E8", "USD 1250")]，按工作表序号与单元格写入（全部按文本写入）。"""
    wb = openpyxl.Workbook()
    sheets: dict[int, object] = {1: wb.active}
    for loc, text in cells:
        _, s, ref = loc.split("-", 2)
        n = int(s[1:])
        while n not in sheets:
            sheets[len(sheets) + 1] = wb.create_sheet(f"Sheet{len(sheets) + 1}")
        sheets[n][ref] = text
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def golden_samples() -> list[Sample]:
    """脱敏回归集：普通样本与 OCR 文字块样本写成 DOCX（每行一段），单元格样本写成 XLSX。"""
    out = []
    for txt in sorted(GOLDEN.glob("*.txt")):
        ann = yaml.safe_load(txt.with_suffix(".yaml").read_text(encoding="utf-8"))
        lines = [line for line in txt.read_text(encoding="utf-8").splitlines() if line.strip()]
        if ann.get("loc_prefixed"):
            pairs = [tuple(line.split(" ", 1)) for line in lines]
            if all("-S" in loc for loc, _ in pairs):
                name, data = f"{txt.stem}.xlsx", _xlsx_bytes(pairs)
            else:  # OCR 文字块：只保留文字，按段落写入
                name, data = f"{txt.stem}.docx", _docx_bytes([t for _, t in pairs])
        else:
            name, data = f"{txt.stem}.docx", _docx_bytes(lines)
        out.append(Sample(name, data, ann["stance"], [i["text"] for i in ann["must_replace"]],
                          list(ann.get("must_suspect", []))))
    return out


def fixture_samples() -> list[Sample]:
    """docx / pdf / xlsx / images 下的全部样本文件（预期解析失败的除外）。"""
    out = []
    for d in FIXTURE_DIRS:
        manifest = yaml.safe_load((FIXTURES / d / "manifest.yaml").read_text(encoding="utf-8"))
        by_file = {v["file"]: v for v in (manifest.get("samples") or {}).values() if isinstance(v, dict) and "file" in v}
        for f in sorted((FIXTURES / d).iterdir()):
            if f.suffix.lower() not in (".docx", ".pdf", ".xlsx", ".png", ".jpg") or f.name in PARSE_FAIL_EXPECTED:
                continue
            ann = by_file.get(f.name, {})
            texts = [i["text"] for i in ann.get("must_replace") or [] if i["text"] not in NOT_IN_SPEC]
            # 标注为必须替换、但自动识别不了而被列为疑似遗漏的（如扫描件中标签识别错的账号），用户会选“替换”
            out.append(Sample(f.name, f.read_bytes(), "sales", texts, texts))
    return out


# ---------------------------------------------------------------- 接口流程


def url(task_id, suffix: str) -> str:
    return f"{API_PREFIX}/tasks/{task_id}/{suffix}"


def upload_and_parse(c, h, db_sessionmaker, sample: Sample, *, source: str = "outgoing") -> uuid.UUID:
    r = c.post(f"{API_PREFIX}/tasks", headers=h,
               data={"stance": sample.stance, "source": source, "policy": "ai", "main_index": "0"},
               files=[("files", (sample.name, sample.data, "application/octet-stream"))])
    assert r.status_code == 201, r.text
    task_id = uuid.UUID(r.json()["id"])
    with db_sessionmaker() as db:  # 代替 worker 执行 parse
        pipeline.parse_task(db, task_id)
        db.execute(update(Job).where(Job.payload["task_id"].astext == str(task_id)).values(status="done"))
        db.commit()
        status = db.get(Task, task_id).status
    assert status == "awaiting_confirm", f"{sample.name}：{status}"
    return task_id


def _matches(text: str, needles: list[str]) -> bool:
    flat = " ".join(text.split())
    return any(n in flat or flat in n for n in needles)


def settle_suspects(c, h, task_id, must_suspect: list[str]) -> None:
    """must_suspect 选替换（长串数字为“其他”，其余为“公司”），其余疑似遗漏忽略。"""
    for _ in range(200):
        pending = [s for s in c.get(url(task_id, "redaction")).json()["suspects"] if s["decision"] == "pending"]
        if not pending:
            return
        s = pending[0]
        if _matches(s["text"], must_suspect):
            body = {"decision": "replace", "category": "CUSTOM" if s["reason"] == "long_digits" else "ORG"}
        else:
            body = {"decision": "ignore"}
        r = c.post(url(task_id, f"redaction/suspects/{s['id']}"), json=body, headers=h)
        assert r.status_code == 200, r.text
    raise AssertionError("疑似遗漏处理不完")


def ensure_our_party(c, h, task_id) -> None:
    """没有我方时：有签约方就把第一个设为我方；一个也没有（样本里没有公司名）就手动补充一段文字作为公司再设为我方。

    补充的文字取最长的、没有任何替换的正文节点的前 40 个字：太短的词（如“采购合同”）可能也出现在提示词或审核要求里，
    会被出口复检按对照表真实值拦截。
    """
    items = c.get(url(task_id, "parties")).json()["items"]
    if any(p["is_our_party"] for p in items):
        return
    if not items:
        files = c.get(f"{API_PREFIX}/tasks/{task_id}").json()["files"]
        nodes = c.get(url(task_id, f"files/{files[0]['id']}/nodes")).json()["items"]
        taken = {i["loc"] for i in c.get(url(task_id, "redaction")).json()["items"]}
        node = max((n for n in nodes if n["part"] == "body" and n["text"] and n["loc"] not in taken),
                   key=lambda n: len(n["text"].strip()))
        text = node["text"]
        start = len(text) - len(text.lstrip())
        end = start + min(40, len(text.strip()))
        while end < len(text) and (text[end - 1] + text[end]).isascii() and (text[end - 1] + text[end]).isalnum():
            end -= 1  # 不在英文单词中间截断
        r = c.post(url(task_id, "redaction/manual"), headers=h, json={
            "loc": node["loc"], "char_start": start, "char_end": end, "category": "ORG"})
        assert r.status_code == 200, r.text
        items = c.get(url(task_id, "parties")).json()["items"]
    body = {"items": [{"id": p["id"], "is_our_party": i == 0, "country": p["country"], "org_type": p["org_type"]}
                      for i, p in enumerate(items)]}
    r = c.put(url(task_id, "parties"), json=body, headers=h)
    assert r.status_code == 200, r.text


def confirm_and_review(c, h, db_sessionmaker, task_id) -> dict:
    """确认发送并（代替 worker）执行审核，返回轮次详情。"""
    r = c.post(url(task_id, "runs"), headers=h)
    assert r.status_code == 201, r.text
    run_id = r.json()["id"]
    with db_sessionmaker() as db:
        job.run_review(db, uuid.UUID(run_id))
        db.execute(update(Job).where(Job.payload["run_id"].astext == run_id).values(status="done"))
        db.commit()
    r = c.get(f"{API_PREFIX}/runs/{run_id}")
    assert r.status_code == 200, r.text
    return r.json()


def run_sample(c, h, db_sessionmaker, sample: Sample) -> uuid.UUID:
    task_id = upload_and_parse(c, h, db_sessionmaker, sample)
    settle_suspects(c, h, task_id, sample.must_suspect)
    ensure_our_party(c, h, task_id)
    settle_suspects(c, h, task_id, sample.must_suspect)  # 手动补充后重算可能产生新的疑似遗漏
    run = confirm_and_review(c, h, db_sessionmaker, task_id)
    assert run["status"] == "completed", f"{sample.name}：{run['status']} {run.get('error_message')}"
    return task_id


# ---------------------------------------------------------------- 外发记录


@dataclass
class Outbound:
    """一次外发：llm_calls 记录、模拟服务实际收到的文字，以及发送时应当检查的真实值。"""
    call_id: int
    task_id: uuid.UUID
    status: str
    request_text: str  # llm_calls.request_text
    wire_text: str  # 模拟服务实际收到的请求体中全部消息的文字
    real_values: list[str]  # 本任务对照表中的全部真实值
    terms: list[str]  # 发送时已存在的已记住名称和自定义敏感词


def outbound_records(db_sessionmaker, requests) -> list[Outbound]:
    """按发送顺序返回全部外发；每条 llm_calls 记录对应模拟服务收到的一个请求（本流程不重试）。"""
    out = []
    with db_sessionmaker() as db:
        calls = list(db.scalars(select(LlmCall).order_by(LlmCall.id)))
        assert len(calls) == len(requests)
        terms = [(t.created_at, decrypt_text(t.term_enc)) for t in db.scalars(select(RedactionTerm))]
        for call, req in zip(calls, requests):
            body = json.loads(req.content)
            out.append(Outbound(
                call.id, call.task_id, call.status, call.request_text or "",
                "\n".join(m.get("content", "") for m in body.get("messages", [])),
                load_real_values(db, call.task_id),
                [text for created, text in terms if created <= call.created_at],
            ))
    return out
