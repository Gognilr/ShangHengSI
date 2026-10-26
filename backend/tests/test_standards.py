"""T-46 管理、加密、向量选择与角色；只使用虚构DOCX。"""
import io
import math

import pytest
from docx import Document
from sqlalchemy import select,func

from app.security import crypto
from app.db.models import Job,Task,TaskClauseVector
from tests.review_support import env,ready_task,llm  # noqa: F401


def document(count=12):
    doc=Document()
    for i in range(count):
        doc.add_heading(f'{i+1} 示例条款',level=1)
        doc.add_paragraph('双方依约交付示例货物。')
    out=io.BytesIO();doc.save(out);return out.getvalue()


def vector(score=1):
    return [score, math.sqrt(1-score*score)]+[0.0]*1022


@pytest.fixture
def template(db_sessionmaker,make_user,env):
    from app.standards.service import create
    user=make_user('template_boss',role='biz_admin')
    with db_sessionmaker() as db:
        row=create(db,document(),dict(name='示例标准合同',doc_types=['contract'],stances=['sales']),user.id)
        db.commit()
        return row


def test_t46_upload_encrypt_and_job(template,db_sessionmaker):
    from app.db.models import StandardClause
    with db_sessionmaker() as db:
        rows=db.scalars(select(StandardClause)).all()
        assert len(rows)==12 and template.clause_count==12
        assert crypto.decrypt_text(template.name_enc)=='示例标准合同'
        assert '示例标准合同'.encode() not in template.name_enc
        assert all('交付'.encode() not in c.text_enc for c in rows)
        assert db.scalar(select(Job).where(Job.type=='template_embed'))


def test_t46_empty_and_limit(template,db_sessionmaker):
    from app.standards.service import create
    from app.security.errors import AppError
    from app.db.models import StandardTemplate
    with db_sessionmaker() as db:
        with pytest.raises(AppError) as e: create(db,document(0),dict(name='示例',doc_types=['contract'],stances=['sales']),template.created_by)
        assert e.value.status==422
        for i in range(49):
            db.add(StandardTemplate(name_enc=crypto.encrypt_text('虚构'),doc_types=['contract'],stances=['sales'],
                                   content_hash='0'*64,clause_count=1,created_by=template.created_by))
        db.flush()
        with pytest.raises(AppError) as e: create(db,document(1),dict(name='示例',doc_types=['contract'],stances=['sales']),template.created_by)
        assert e.value.status==422


def test_t46_clause_hash_and_safe_audit(template,db_sessionmaker):
    from app.standards.service import update_clause
    from app.db.models import StandardClause,StandardTemplate,AuditLog
    with db_sessionmaker() as db:
        c=db.scalar(select(StandardClause))
        before=db.get(StandardTemplate,template.id).content_hash
        update_clause(db,template.id,c.id,dict(stance='示例标准立场',acceptable='示例范围',escalate='示例上报'),template.created_by)
        db.commit()
        assert db.get(StandardTemplate,template.id).content_hash!=before
        audit=db.scalar(select(AuditLog).where(AuditLog.action=='template_clause_update'))
        assert '示例标准立场' not in str(audit.detail)


@pytest.mark.parametrize('condition',['ready','disabled','type','stance','incomplete','low'])
def test_t46_selection(template,db_sessionmaker,condition):
    from app.standards.service import choose,create
    from app.db.models import StandardClause,StandardTemplate
    with db_sessionmaker() as db:
        t=Task(owner_id=template.created_by,title='示例',stance='sales',source='incoming',policy='ai',status='awaiting_confirm')
        db.add(t);db.flush()
        db.add(TaskClauseVector(task_id=t.id,file_seq=1,first_loc='F1-P1',locs=['F1-P1'],embedding=vector(),model='test'))
        for c in db.scalars(select(StandardClause)): c.embedding=vector(0.65 if condition=='low' else 0.82)
        model=db.get(StandardTemplate,template.id)
        if condition=='disabled':model.enabled=False
        if condition=='type':model.doc_types=['quotation']
        if condition=='stance':model.stances=['purchase']
        if condition=='incomplete':db.scalar(select(StandardClause)).embedding=None
        db.flush()
        if condition=='ready':
            other=create(db,document(1),
                dict(name='低相似度候选',doc_types=['contract'],stances=['sales']),template.created_by)
            db.scalar(select(StandardClause).where(StandardClause.template_id==other.id)).embedding=vector(0.65)
            db.flush()
        selected=choose(db,t,'contract')
        assert (selected is not None)==(condition=='ready')
        if selected is not None:assert selected.id==template.id


@pytest.mark.parametrize('role,expected',[('user',403),('sys_admin',403),('biz_admin',200),('admin',200)])
def test_t46_roles(make_user,login,role,expected,env):
    make_user('templates_role',role=role)
    c,csrf=login('templates_role')
    assert c.get('/api/v1/admin/templates').status_code==expected
    assert c.post('/api/v1/admin/templates',headers={'X-CSRF-Token':csrf},
        data={'name':'虚构范本','doc_types':'["contract"]','stances':'["sales"]'},
        files={'file':('示例.docx',document(1))}).status_code==(201 if expected==200 else expected)


def test_t46_snapshot_modified_deleted_and_no_vector(template,db_sessionmaker):
    from app.db.models import StandardTemplate,ReviewRun
    from app.standards.service import snapshot,view
    with db_sessionmaker() as db:
        task=Task(owner_id=template.created_by,title='示例',stance='sales',source='incoming',policy='ai',status='reviewed')
        db.add(task);db.flush()
        run=ReviewRun(task_id=task.id,run_no=1,status='completed',stance='sales',source='incoming',main_doc_type='contract',
                      created_by=template.created_by,template_id=template.id,template_name_snapshot_enc=template.name_enc,template_hash=template.content_hash)
        db.add(run);db.commit()
        row=db.get(StandardTemplate,template.id)
        assert 'embedding' not in str(view(db,row,True))
        row.content_hash='a'*64;db.commit()
        assert snapshot(db,run)==dict(name='示例标准合同',status='modified')
        db.delete(row);db.commit();db.refresh(run)
        assert run.template_id is None and snapshot(db,run)['status']=='deleted'


def test_t46_rotation_columns():
    from app.db.base import Base
    columns={(t.name,c.name) for t in Base.metadata.tables.values() for c in t.columns if c.name.endswith('_enc')}
    assert {('standard_templates','name_enc'),('review_runs','template_name_snapshot_enc'),('law_import_files','data_enc')} <= columns
    assert {('standard_clauses',name+'_enc') for name in ('title','text','stance','acceptable','escalate')} <= columns


def test_t46_migration_upgrade_downgrade():
    import os,uuid,psycopg
    from alembic import command
    from tests.test_migrations import _alembic,_url_for
    url=os.environ.get('TEST_DATABASE_URL')
    if not url:pytest.skip('无测试库')
    name='aireview_test_'+uuid.uuid4().hex[:10]
    with psycopg.connect(url,autocommit=True) as admin:admin.execute(f'CREATE DATABASE "{name}"')
    try:
        cfg=_alembic(_url_for(name));command.upgrade(cfg,'0022')
        with psycopg.connect(_url_for(name),autocommit=True) as db:
            assert db.execute("SELECT character_maximum_length FROM information_schema.columns WHERE table_name='jobs' AND column_name='type'").fetchone()==(20,)
            db.execute("INSERT INTO jobs(type,payload,created_at) VALUES('template_embed','{}',now())")
        command.downgrade(cfg,'0021')
        with psycopg.connect(_url_for(name),autocommit=True) as db:
            assert db.execute("SELECT to_regclass('standard_templates')").fetchone()==(None,)
            assert db.execute("SELECT 1 FROM jobs WHERE type='template_embed'").fetchone() is None
            assert db.execute("SELECT character_maximum_length FROM information_schema.columns WHERE table_name='jobs' AND column_name='type'").fetchone()==(12,)
        command.upgrade(cfg,'head')
    finally:
        with psycopg.connect(url,autocommit=True) as admin:admin.execute(f'DROP DATABASE "{name}" WITH (FORCE)')


@pytest.mark.parametrize('use_template',[False,True])
def test_t46_run_selection_and_worker(db_sessionmaker,make_user,env,ready_task,llm,monkeypatch,use_template):
    from app.standards import service
    from app.embedding import engine
    from app.redaction import ner
    from app.review import service as reviews,job
    from app.db.models import LocalFinding,ReviewRun,Finding
    monkeypatch.setattr(ner,'recognize',lambda text:[])
    monkeypatch.setattr(engine,'embed',lambda texts,kind:[vector() for _ in texts])
    user=make_user('run_template_owner')
    task=ready_task(user)
    requests=llm('normal')
    with db_sessionmaker() as db:
        template=service.create(db,document(1),dict(name='虚构对照范本',doc_types=['contract'],stances=['sales']),user.id)
        for c in service.clauses(db,template.id):c.embedding=vector()
        db.add(LocalFinding(task_id=task.id,check_code='L-12',severity='low',locs=[],detail_enc=crypto.encrypt_text('旧提示'),ai_summary='旧提示'))
        db.commit()
        run=reviews.create_run(db,db.get(Task,task.id),user,None,use_template=use_template)
        assert (run.template_id is not None)==use_template
        job.run_review(db,run.id)
        db.refresh(run)
        assert run.status=='completed'
        assert db.query(LocalFinding).filter_by(check_code='L-12').count()==0
        assert (reviews.run_detail(db,run).template is not None)==use_template
    import json
    sent=json.loads(requests[0].content)['messages'][1]['content']
    assert ('【公司范本对照】' in sent)==use_template
    assert '旧提示' not in sent
