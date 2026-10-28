"""T-51：虚构知识库、模拟向量、真实权限与加密边界。"""
import io
import json
import uuid

import pytest
from docx import Document
from sqlalchemy import select, func

from app.admin.settings_store import set_setting
from app.db.models import AuditLog, Job, Task
from app.security import crypto
from tests.review_support import env  # noqa: F401
from tests.test_assistant_chat import conversation  # noqa: F401


def vector():
    return [1.0] + [0.0] * 1023


def document():
    doc = Document()
    doc.add_heading('虚构采购制度', level=1)
    for _ in range(5):
        doc.add_paragraph('应当保留采购审批记录。' * 30)
    doc.add_heading('虚构交付制度', level=1)
    doc.add_paragraph('应当核对交付日期。')
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip([c for row in table.rows for c in row.cells], ['事项', '要求', '审批', '留档']):
        cell.text = value
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.mark.parametrize('role,expected', [('user', 403), ('sys_admin', 403), ('biz_admin', 200), ('admin', 200)])
def test_t51_1_all_routes_roles(env, make_user, login, role, expected):
    user = make_user('kb_role', role=role)
    c, csrf = login(user.username)
    h = {'X-CSRF-Token': csrf}
    base = '/api/v1/admin/kb'
    assert c.get(base).status_code == expected
    r = c.post(base, headers=h, files={'file': ('虚构制度.docx', document())})
    assert r.status_code == (201 if expected == 200 else 403)
    kid = r.json()['id'] if expected == 200 else str(uuid.uuid4())
    for response in (c.get(f'{base}/{kid}/chunks'), c.patch(f'{base}/{kid}', headers=h, json={'enabled': False}),
                     c.delete(f'{base}/{kid}', headers=h)):
        assert response.status_code == expected


def test_t51_2_upload_parse_encrypt_duplicate_preview_and_audit(env, make_user, login, db_sessionmaker, monkeypatch, caplog):
    from app.db.models import KbDoc, KbChunk
    from app.kb.job import handle_import
    from app.embedding import engine
    monkeypatch.setattr(engine, 'embed', lambda texts, kind: [vector() for _ in texts])
    user = make_user('kb_boss', role='biz_admin')
    c, csrf = login(user.username)
    h = {'X-CSRF-Token': csrf}
    base = '/api/v1/admin/kb'
    data = document()
    with caplog.at_level('INFO'):
        r = c.post(base, headers=h, files={'file': ('虚构秘密制度.docx', data)})
        assert r.status_code == 201, r.text
        kid = uuid.UUID(r.json()['id'])
        assert c.post(base, headers=h, files={'file': ('另一名称.docx', data)}).status_code == 409
        with db_sessionmaker() as db:
            job = db.scalar(select(Job).where(Job.type == 'kb_import'))
            import base64
            encrypted = base64.b64decode(job.payload['data_enc'])
            assert not encrypted.startswith(b'PK') and crypto.decrypt(encrypted).startswith(b'PK')
            handle_import(db, job.payload)
        preview = c.get(f'{base}/{kid}/chunks')
    assert preview.status_code == 200
    items = preview.json()['items']
    assert 3 <= len(items) <= 6 and '事项：审批' in str(items) and '要求：留档' in str(items)
    assert any(item['heading'] == '虚构交付制度' for item in items)
    assert 'embedding' not in preview.text and 'embedding' not in c.get(base).text
    with db_sessionmaker() as db:
        row = db.get(KbDoc, kid)
        assert row.status == 'ready' and row.chunk_count == len(items)
        assert crypto.decrypt_text(row.name_enc) == '虚构秘密制度.docx'
        assert '虚构秘密'.encode() not in row.name_enc
        chunks = db.scalars(select(KbChunk)).all()
        assert all('采购审批'.encode() not in item.text_enc and item.embedding is not None for item in chunks)
        audit = db.scalar(select(AuditLog).where(AuditLog.action == 'kb_upload'))
        assert audit and '制度' not in str(audit.detail)
        assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == 'kb_upload')) == 1
    for secret in ('虚构秘密制度', '采购审批记录', '虚构交付制度'):
        assert secret not in caplog.text


def test_t51_3_disabled_no_engine(conversation, db_sessionmaker, monkeypatch):
    from app.kb.retrieve import retrieve
    from app.embedding import engine
    monkeypatch.setattr(engine, 'embed', lambda *args: pytest.fail('关闭时不应调用向量引擎'))
    with db_sessionmaker() as db:
        task = db.get(Task, uuid.UUID(conversation[2]))
        assert retrieve(db, task, '查询原文') == ([], 0)


def test_t51_4_5_6_redact_discard_and_enabled_only(conversation, db_sessionmaker, monkeypatch, caplog):
    from app.db.models import KbDoc, KbChunk
    from app.kb.retrieve import retrieve
    from app.embedding import engine
    from app.redaction import ner
    from tests.test_assistant_chat import REAL
    monkeypatch.setattr(engine, 'embed', lambda texts, kind: [vector() for _ in texts])
    # 金额/电话由规则识别，疑似项通过与真实 NER 相同的消息脱敏结果注入。
    with db_sessionmaker() as db:
        set_setting(db, 'assistant.kb_enabled', True, None)
        task = db.get(Task, uuid.UUID(conversation[2]))
        for i, (enabled, status, text) in enumerate([(True, 'ready', REAL + '；金额：人民币12000元；电话：13900139000'),
                (True, 'ready', '联系星海公司核对'), (False, 'ready', '禁用片段'), (True, 'failed', '失败片段')]):
            doc = KbDoc(name_enc=crypto.encrypt_text('不可外发名称'), status=status, enabled=enabled,
                        content_hash=str(i) * 64, created_by=task.owner_id)
            db.add(doc); db.flush()
            db.add(KbChunk(doc_id=doc.id, seq=1, text_enc=crypto.encrypt_text(text), embedding=vector()))
        db.commit()
        with caplog.at_level('INFO'):
            items, dropped = retrieve(db, task, '虚构查询秘密')
        assert len(items) == 1 and dropped == 1
        rendered = str(items)
        assert '[[ORG_1]]' in rendered and '[[AMT_' in rendered and '[[TEL_' in rendered
        for secret in (REAL, '12000', '13900139000', '禁用片段', '失败片段', '不可外发名称'):
            assert secret not in rendered
        assert '虚构查询秘密' not in caplog.text
        db.commit()
        for doc in db.scalars(select(KbDoc)):
            db.delete(doc)
        db.commit()
        assert retrieve(db, task, '查询') == ([], 0)


def test_t51_medium_person_dropped_and_recheck_failure(conversation, db_sessionmaker, monkeypatch):
    from app.kb.retrieve import protect
    from app.redaction import ner
    import app.kb.retrieve as module
    from app.redaction.recheck import BlockReason
    monkeypatch.setattr(ner, 'recognize', lambda text: [ner.Entity(2, 4, 'PERSON', 0.70)])
    with db_sessionmaker() as db:
        task = db.get(Task, uuid.UUID(conversation[2]))
        assert protect(db, task, '联系林澄核对') is None
        monkeypatch.setattr(ner, 'recognize', lambda text: [])
        monkeypatch.setattr(module, 'recheck', lambda *args, **kwargs: [BlockReason('MAP', 0)])
        assert protect(db, task, '普通制度条款') is None


def test_t51_kb_limit_score_and_budget(conversation, db_sessionmaker, monkeypatch):
    from app.db.models import KbDoc, KbChunk
    from app.kb.retrieve import retrieve
    from app.embedding import engine
    from tests.test_standards import vector as scored
    monkeypatch.setattr(engine, 'embed', lambda texts, kind: [vector() for _ in texts])
    with db_sessionmaker() as db:
        task = db.get(Task, uuid.UUID(conversation[2]))
        set_setting(db, 'assistant.kb_enabled', True, None)
        doc = KbDoc(name_enc=crypto.encrypt_text('虚构'), status='ready', content_hash='8'*64, created_by=task.owner_id)
        db.add(doc); db.flush()
        for n in range(1, 8):
            db.add(KbChunk(doc_id=doc.id, seq=n, text_enc=crypto.encrypt_text('核对交付物。'*140),
                           embedding=scored(0.9 if n < 7 else 0.3)))
        db.commit()
        items, dropped = retrieve(db, task, '问题')
        assert 1 <= len(items) <= 5 and sum(len(item['text']) for item in items) <= 4000 and dropped == 0


def test_kb_failure_cleanup_delete_and_settings(env, make_user, login, db_sessionmaker, monkeypatch):
    from app.kb.job import handle_import
    import app.kb.job as module
    from app.db.models import KbDoc, KbChunk
    user = make_user('kb_failure', role='biz_admin')
    c, csrf = login(user.username)
    h = {'X-CSRF-Token': csrf}
    base = '/api/v1/admin/kb'
    assert c.patch(base+'/settings', headers=h, json={'enabled': True}).json() == {'enabled': True}
    r = c.post(base, headers=h, files={'file': ('虚构制度.docx', document())})
    kid = uuid.UUID(r.json()['id'])
    monkeypatch.setattr(module, '_parse_bytes', lambda *args: (_ for _ in ()).throw(ValueError('虚构秘密异常原文')))
    with db_sessionmaker() as db:
        job = db.scalar(select(Job).where(Job.type == 'kb_import'))
        with pytest.raises(ValueError):
            handle_import(db, job.payload)
        assert db.get(KbDoc, kid).status == 'failed'
        assert '秘密' not in db.get(KbDoc, kid).error_message
        db.refresh(job)
        assert job.payload == {'doc_id': str(kid)}
    assert c.delete(f'{base}/{kid}', headers=h).status_code == 200
    assert c.get(f'{base}/{kid}/chunks').status_code == 404
    with db_sessionmaker() as db:
        assert db.scalar(select(func.count()).select_from(KbChunk)) == 0


def test_kb_old_conversion_and_pdf_ocr_share_parser(monkeypatch):
    from app.kb.job import _parse_bytes
    from app.parsing import legacy, pdf
    from app.parsing.model import ParsedFile
    from app import worker
    from types import SimpleNamespace
    from contextlib import contextmanager
    calls = []
    monkeypatch.setattr(legacy, 'convert', lambda data, ext: calls.append(ext) or document())
    parsed, _ = _parse_bytes('doc', b'fictional', 1, 0)
    assert parsed.nodes and calls == ['doc']
    @contextmanager
    def locked():
        calls.append('lock')
        yield
    monkeypatch.setattr(worker, 'ocr_slot', locked)
    monkeypatch.setattr(pdf, 'inspect_pdf', lambda data: SimpleNamespace(scanned_pages=[1]))
    monkeypatch.setattr(pdf, 'parse_pdf', lambda data, seq: ParsedFile([], 1, False, [], 1))
    _parse_bytes('pdf', b'fictional', 1, 0)
    assert calls == ['doc', 'lock']
