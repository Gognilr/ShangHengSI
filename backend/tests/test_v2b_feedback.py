"""T-53：真实反馈对应虚构用例，先写测试再实现。"""
import io

import pytest
from docx import Document
from sqlalchemy import delete, select

from app.redaction import detect, ner, parties, suspects
from app.redaction.engine import NodeIn, Term, compute
from app.db.models import RedactionSuspect
from tests.review_support import env, ready_task  # noqa: F401
from tests.test_user_note_redaction import alice, task_url  # noqa: F401


@pytest.fixture(autouse=True)
def offline_ner(monkeypatch):
    monkeypatch.setattr(ner, 'recognize', lambda text: [])


def test_t53_02_company_boundaries():
    text = '我方澄川虚构碳材料股份有限公司，现指定青岚虚构国际贸易有限公司为代理'
    spans = parties.find_other_org_names(text)
    assert [text[a:b] for a, b in spans] == ['澄川虚构碳材料股份有限公司', '青岚虚构国际贸易有限公司']
    result = compute([NodeIn('F1-P1', text)], 'other')
    assert '我方' in result.redacted['F1-P1'] and '现指定' in result.redacted['F1-P1']


@pytest.mark.parametrize('name', ['现代虚构重工有限公司', '特变虚构电工有限公司'])
def test_t53_03_preserve_company_start(name):
    assert [(a, b) for a, b in parties.find_other_org_names(name)] == [(0, len(name))]


def test_t53_04_existing_learned_name():
    name = '我方虚构科技有限公司'
    result = compute([NodeIn('F1-P1', name)], 'other', terms=[Term(text=name, kind='learned', category='ORG')])
    assert result.items and any(i.char_start == 0 and i.char_end == len(name) for i in result.items)


@pytest.mark.parametrize('text,name', [('尊敬的柯墨兰先生：', '柯墨兰'),
    ('请欧阳明月女士确认', '欧阳明月'), ('叶先生', '叶先生'),
    ('DEAR MR.KE MOLAN:', 'KE MOLAN'), ('Dear Mrs Lorna Fictive,', 'Lorna Fictive'),
    ('Attn: MISS JANE FICTIVE', 'JANE FICTIVE'), ('Contact Name: Lorna Fictive', 'Lorna Fictive')])
def test_t53_05_07_09_persons(text, name):
    hits = detect.find_persons(text)
    assert [text[h.start:h.end] for h in hits] == [name]


@pytest.mark.parametrize('text', ['各位先生', '女士们', 'Dr Street 12', 'Main Dr. ',
    'Product Name: Calcined Needle Coke', 'Item Name: Graphite Electrode'])
def test_t53_06_08_09_non_persons(text):
    assert detect.find_persons(text) == []


def test_t53_09_company_name_unchanged():
    text = 'Company Name: Northwind Fictive'
    assert [text[h.start:h.end] for h in detect.find_persons(text)] == ['Northwind Fictive']


@pytest.mark.usefixtures('env')
@pytest.mark.parametrize('decision', ['ignore', 'replace'])
def test_t53_10_group_and_decide(alice, ready_task, db_sessionmaker, decision):
    user, client, headers = alice
    doc = Document()
    for _ in range(5):
        doc.add_paragraph('Velmora Refinery')
    stream = io.BytesIO()
    doc.save(stream)
    task = ready_task(user, files=[('虚构分组.docx', stream.getvalue())])
    with db_sessionmaker() as db:
        db.execute(delete(RedactionSuspect).where(RedactionSuspect.task_id == task.id))
        for i in range(1, 6):
            db.add(RedactionSuspect(task_id=task.id, loc=f'F1-P{i}', char_start=0, char_end=16,
                reason='capitalized' if i % 2 else 'ner', decision='pending'))
        db.commit()
    response = client.get(task_url(task, '/redaction'), headers=headers)
    assert response.status_code == 200
    rows = response.json()['suspects']
    assert len(rows) == 1 and rows[0]['occurrences'] == 5
    assert rows[0]['loc'] == 'F1-P1' and rows[0]['reason'] == 'capitalized'
    response = client.post(task_url(task, f"/redaction/suspects/{rows[0]['id']}"), headers=headers,
        json={'decision': decision, 'category': 'ORG' if decision == 'replace' else None})
    assert response.status_code == 200, response.text
    assert not [s for s in response.json()['suspects'] if s['decision'] == 'pending']
    if decision == 'replace':
        assert len(response.json()['items']) == 5
    else:
        with db_sessionmaker() as db:
            assert len(list(db.scalars(select(RedactionSuspect).where(
                RedactionSuspect.task_id == task.id, RedactionSuspect.decision == 'ignore')))) == 5


@pytest.mark.usefixtures('env')
def test_t53_11_two_groups(alice, ready_task):
    user, client, headers = alice
    doc = Document()
    for text in ['Velmora Refinery', 'Hardgrave Office'] * 3:
        doc.add_paragraph(text)
    stream = io.BytesIO()
    doc.save(stream)
    task = ready_task(user, files=[('虚构两组.docx', stream.getvalue())])
    rows = client.get(task_url(task, '/redaction'), headers=headers).json()['suspects']
    assert len(rows) == 2 and all(r['occurrences'] == 3 for r in rows)


@pytest.mark.parametrize('text', ['approved suppliers', 'ASTM D4052', 'GB/T 24526',
    'T/ZGTS 002-2022', 'ASTM', 'We', 'The inspector'])
@pytest.mark.parametrize('score', [0.5, 0.95])
def test_t53_12_ner_filter(text, score):
    assert not ner.keep_entity(text, ner.Entity(0, len(text), 'ORG', score))


@pytest.mark.parametrize('text', ['Velmora Refinery', 'CCIC'])
@pytest.mark.parametrize('score', [0.5, 0.95])
def test_t53_13_ner_kept(text, score):
    assert ner.keep_entity(text, ner.Entity(0, len(text), 'ORG', score))


def test_t53_14_15_lowercase_evidence(monkeypatch):
    # V2Q-169：仅本用例隔离后续 B-14 词表，保留 B-12 的原输入及全部断言。
    monkeypatch.setattr(suspects, 'en_words', lambda: frozenset(), raising=False)
    nodes = [NodeIn('F1-P1', 'B. Feedstock Base'), NodeIn('F1-P2', 'feedstock base')]
    assert not compute(nodes, 'other').suspects
    text = 'B. Velmora Base'
    result = compute([NodeIn('F1-P1', text)], 'other')
    assert [text[s.char_start:s.char_end] for s in result.suspects] == ['Velmora Base']
    text = 'Kinematic Viscosity'
    assert compute([NodeIn('F1-P1', text)], 'other').suspects
    assert not compute([NodeIn('F1-P1', text), NodeIn('F1-P2', text.lower())], 'other').suspects


def test_t53_16_chat_evidence():
    from app.chat.service import _redact
    assert not _redact('Bulk Density 偏低怎么办，bulk density 指什么？', {}, []).suspects


@pytest.mark.parametrize('text', ['Calcined Needle Coke', 'Product Name: Calcined Needle Coke'])
def test_v2q167_standard_product_preserved(text):
    result = compute([NodeIn('F1-P1', text)], 'other')
    assert result.redacted['F1-P1'] == text
    start = text.index('Calcined Needle Coke')
    assert not result.items
    assert not any(s.char_start < len(text) and s.char_end > start for s in result.suspects)


def test_v2q167_product_label_does_not_exempt_proper_name():
    text = 'Product Name: Velmora Needle Coke'
    result = compute([NodeIn('F1-P1', text)], 'other')
    start = text.index('Velmora')
    assert any(s.char_start <= start and s.char_end >= start + len('Velmora') for s in result.suspects)


@pytest.mark.parametrize('english', ['Chengchuan Fictive Carbon Materials Co., Ltd.'])
def test_t53_17_bilingual_same_code_and_restore(english):
    from app.redaction.restore import restore
    chinese = '澄川虚构碳材料股份有限公司'
    nodes = [NodeIn('F1-P1', chinese), NodeIn('F1-P2', english)]
    result = compute(nodes, 'other')
    assert len(result.parties) == 1
    assert result.parties[0].country == '中国'
    assert result.redacted['F1-P1'] == result.redacted['F1-P2']
    code = result.redacted['F1-P1']
    assert result.mapping[code] == chinese
    assert english in [v for k, v in result.mapping.items() if k.startswith(code + '#')]
    assert restore(result.redacted['F1-P1'], result.mapping) == chinese
    # 同一代号的别名保留原英文写法；按来源选择该写法后可逐字重放。
    alias = next(v for k, v in result.mapping.items() if k.startswith(code + '#') and v == english)
    assert restore(result.redacted['F1-P2'], {code: alias}) == english


@pytest.mark.parametrize('english', ['Qingdao Qinglan Fictive International Trade Co., Ltd.',
    'Qinglan Fictive Trade Co., Ltd.'])
def test_t53_18_bilingual_geography(english):
    result = compute([NodeIn('F1-P1', '青岚虚构国际贸易有限公司'), NodeIn('F1-P2', english)], 'other')
    assert len(result.parties) == 1 and result.redacted['F1-P1'] == result.redacted['F1-P2']


def test_t53_19_bilingual_ambiguity():
    result = compute([NodeIn('F1-P1', '华信虚构科技有限公司'), NodeIn('F1-P2', '华新虚构科技有限公司'),
        NodeIn('F1-P3', 'Huaxin Fictive Co., Ltd.')], 'other')
    assert len(result.parties) == 3
    assert len(set(result.redacted.values())) == 3


def test_t53_20_non_pinyin_company():
    from app.redaction.parties import bilingual_links
    assert not bilingual_links(['澄川虚构科技有限公司', 'Northwind Fictive International Limited'])


def test_t53_21_chat_bilingual_existing_code():
    from app.chat.service import _redact
    chinese, english = '澄川虚构碳材料股份有限公司', 'Chengchuan Fictive Carbon Materials Co., Ltd.'
    result = _redact(english, {'公司_A': chinese}, [])
    assert result.text == '公司_A'
    assert english in result.new_entries.values()
    assert all(k.startswith('公司_A#') for k in result.new_entries)


def test_t53_13_pinyin_resource():
    from app.redaction.dictionaries import hanzi_pinyin
    readings = hanzi_pinyin()
    assert len(readings) == 6763
    assert {'cheng', 'deng'} <= set(readings['澄'])
    assert {'chong', 'zhong'} <= set(readings['重'])
    assert all(value and all(p.isascii() and p.isalpha() and p.islower() for p in value)
               for value in readings.values())


def test_t53_13_ambiguity_in_reverse_direction():
    from app.redaction.parties import bilingual_links
    assert not bilingual_links(['澄川虚构科技有限公司', 'Chengchuan Fictive Co., Ltd.',
        'Chengchuan Other Co., Ltd.'])


@pytest.mark.usefixtures('env')
def test_t53_13_alias_not_learned(alice, ready_task, db_sessionmaker):
    from app.db.models import RedactionTerm, Task
    from app.redaction.engine import learn_confirmed_names
    from app.security.crypto import decrypt_text
    user = alice[0]
    chinese, english = '澄川虚构碳材料股份有限公司', 'Chengchuan Fictive Carbon Materials Co., Ltd.'
    doc = Document()
    doc.add_paragraph(chinese)
    doc.add_paragraph(english)
    stream = io.BytesIO()
    doc.save(stream)
    task = ready_task(user, files=[('虚构双语.docx', stream.getvalue())])
    with db_sessionmaker() as db:
        learn_confirmed_names(db, db.get(Task, task.id), user.id)
        names = [decrypt_text(r.term_enc) for r in db.scalars(select(RedactionTerm))]
        assert chinese in names and english not in names


@pytest.mark.parametrize('text', ['Kinematic Viscosity', 'Document Title'])
def test_t53_22_common_english_words(text):
    assert not compute([NodeIn('F1-P1', text)], 'other').suspects


@pytest.mark.parametrize('text', ['Velmora Refinery', 'Hardgrave Machine Method'])
def test_t53_23_uncommon_name_still_suspect(text):
    result = compute([NodeIn('F1-P1', text)], 'other')
    assert any(s.char_start == 0 and s.char_end == len(text) for s in result.suspects)


def test_t53_24_spec_capitalized_groups(monkeypatch):
    from pathlib import Path
    import yaml
    root = Path(__file__).parent / 'fixtures/redaction_golden'
    lines = (root / 'en_spec_suspects.txt').read_text(encoding='utf-8').splitlines()
    annotation = yaml.safe_load((root / 'en_spec_suspects.yaml').read_text(encoding='utf-8'))
    assert len(lines) >= 200
    # T-24 的 NER 组数需如实报告，使用已选离线模型，不用分数模拟。
    monkeypatch.setattr(ner, 'recognize', lambda text: [e for e in ner.raw_recognize(text) if ner.keep_entity(text, e)])
    nodes = [NodeIn(f'F1-P{i}', line) for i, line in enumerate(lines, 1) if line.strip()]
    result = compute(nodes, annotation['stance'])
    texts = {n.loc: n.text for n in nodes}
    groups = {}
    for s in result.suspects:
        key = parties.normalize(texts[s.loc][s.char_start:s.char_end])
        previous = groups.get(key)
        if previous is None or suspects._PRIORITY.get(s.reason, 5) < suspects._PRIORITY.get(previous, 5):
            groups[key] = s.reason
    capitalized = sum(reason == 'capitalized' for reason in groups.values())
    model = sum(reason == 'ner' for reason in groups.values())
    print('T-24 capitalized groups:', capitalized, 'NER groups:', model, 'total groups:', len(groups))
    assert capitalized <= 20
    for needle in annotation['must_suspect']:
        for n in nodes:
            start = n.text.find(needle)
            if start >= 0:
                assert any(s.loc == n.loc and s.char_start <= start and s.char_end >= start + len(needle)
                           for s in result.suspects), needle


def test_t53_14_word_resource():
    from pathlib import Path
    from app.redaction.dictionaries import en_words
    words = en_words()
    assert {'kinematic', 'viscosity', 'document', 'title', 'global'} <= words
    assert not {'velmora', 'hardgrave'} & words
    assert (Path(__file__).parents[1] / 'resources/en_words.txt').stat().st_size <= 1_000_000
    assert all(w.isascii() and w.isalpha() and w.islower() for w in words)
