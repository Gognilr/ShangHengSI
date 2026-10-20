"""生成 M6 的虚构金额样本；只在 Docker 测试镜像中运行。金额用字符串，标注独立手写。"""
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

import docx
import openpyxl
import yaml
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas

OUT = Path(__file__).resolve().parent
FIXED = datetime(2026, 1, 1)
HEADER = ["项目", "数量", "单价 USD", "金额 USD"]


def frozen_zip(data):
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            entry = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            content = src.read(name)
            if name == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                                 rb"\g<1>2026-01-01T00:00:00Z\g<2>", content)
            dst.writestr(entry, content)
    return stream.getvalue()


def xlsx(name, rows):
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = FIXED
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    (OUT / name).write_bytes(frozen_zip(buf.getvalue()))


def word(name, rows):
    document = docx.Document()
    document.core_properties.created = document.core_properties.modified = FIXED
    document.add_paragraph("虚构报价单，用于自动化校验。")
    table = document.add_table(rows=0, cols=len(rows[0]))
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    buf = io.BytesIO()
    document.save(buf)
    (OUT / name).write_bytes(frozen_zip(buf.getvalue()))


def pdf(name):
    canvas = Canvas(str(OUT / name), pagesize=A4, invariant=1)
    canvas.setFont("Helvetica", 11)
    canvas.drawString(50, 780, "Fictional quotation for local checks only")
    rows = [["Item", "Qty", "Unit Price USD", "Amount USD"],
            ["Fictional product", "100", "520", "52400"], ["Total", "", "", "52400"]]
    xs, top, height = [50, 210, 280, 410, 550], 740, 26
    for i in range(len(rows) + 1):
        canvas.line(xs[0], top - i * height, xs[-1], top - i * height)
    for x in xs:
        canvas.line(x, top, x, top - len(rows) * height)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            canvas.drawString(xs[j] + 5, top - i * height - 17, text)
    canvas.save()


def contract(name, text):
    document = docx.Document()
    document.core_properties.created = document.core_properties.modified = FIXED
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    (OUT / name).write_bytes(frozen_zip(buf.getvalue()))


def main():
    xlsx("row_error.xlsx", [HEADER, ["虚构焦炭甲", "100", "520", "52400"], ["合计", "", "", "52400"]])
    xlsx("total_error.xlsx", [HEADER, ["虚构焦炭甲", "2", "5000", "10000"],
                              ["虚构焦炭乙", "1", "5000", "5000"], ["合计", "", "", "14000"]])
    xlsx("tax_error.xlsx", [["数量", "单价", "不含税金额", "税率", "税额", "含税金额"],
                            ["10", "1000", "10000", "13%", "1200", "11300"]])
    xlsx("unreadable.xlsx", [HEADER, ["虚构焦炭甲", "待定", "520", "无法识别"], ["合计", "", "", "0"]])
    xlsx("formula_uncached.xlsx", [HEADER, ["虚构焦炭甲", "1", "100", "=B2*C2"], ["合计", "", "", "100"]])
    word("valid_quote.docx", [HEADER, ["虚构焦炭甲", "2", "5", "10"],
                              ["虚构焦炭乙", "3", "7.5", "22.5"], ["合计", "", "", "32.5"]])
    pdf("row_error.pdf")
    contract("uppercase_error.docx", "合同金额：人民币850000元，大写捌拾万元整。")
    contract("uppercase_ok.docx", "合同金额：人民币80万元，大写捌拾万元整。")
    contract("contract_total.docx", "合同总价：人民币800000元。")
    for currency in ("CNY", "USD"):
        xlsx(f"quote_{currency.lower()}.xlsx", [["数量", "单价", f"金额 {currency}"],
                                               ["1", "850000", "850000"], ["合计", "", "850000"]])
    # 标注不通过被测校验代码推算；金额和位置均为明确的独立预期。
    samples = {
        "row_error.xlsx": [("L-01", ["F1-S1-D2", "F1-S1-B2", "F1-S1-C2"])],
        "total_error.xlsx": [("L-02", ["F1-S1-D4"])],
        "tax_error.xlsx": [("L-03", ["F1-S1-E2"]), ("L-03", ["F1-S1-F2"])],
        "unreadable.xlsx": [("L-INFO", ["F1-S1-A1"])],
        "formula_uncached.xlsx": [("L-INFO", ["F1-S1-A1"]), ("L-INFO", [])],
        "valid_quote.docx": [],
        "row_error.pdf": [("L-01", ["F1-PG1-T1-R2-C4", "F1-PG1-T1-R2-C2", "F1-PG1-T1-R2-C3"])],
        "uppercase_error.docx": [("L-04", ["F1-P1"])],
        "uppercase_ok.docx": [], "contract_total.docx": [], "quote_cny.xlsx": [], "quote_usd.xlsx": [],
    }
    manifest = {"description": "全部为虚构样本；校验位置假定 file_seq=1；M6-01 覆盖 L-01～L-03、L-INFO。",
                "samples": [{"file": name, "doc_type": "contract" if name.startswith(("uppercase", "contract")) else "quotation",
                             "expected": [{"code": code, "locs": locs} for code, locs in expected]}
                            for name, expected in samples.items()]}
    manifest["consistency_cases"] = [
        {"files": ["contract_total.docx", f"quote_{currency}.xlsx"],
         "expected": [{"code": code, "locs": ["F1-P1", "F2-S1-C3"]}]}
        for currency, code in (("cny", "L-05"), ("usd", "L-06"))]
    (OUT / "manifest.yaml").write_text(yaml.safe_dump(manifest, allow_unicode=True, sort_keys=False), encoding="utf-8")


if __name__ == "__main__":
    main()
