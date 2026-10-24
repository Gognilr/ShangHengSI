"""V2-B-03 设计 3.6 测试先行。"""
import pytest
import io
from docx import Document
from sqlalchemy import select
from app.db.models import RedactionTerm
from app.redaction.engine import learn_confirmed_names
from app.security.crypto import decrypt_text
from tests.review_support import env, ready_task  # noqa: F401
from tests.test_redaction_api import alice  # noqa: F401
from app.chat.service import _redact
from app.redaction.engine import NodeIn, compute
from tests.test_redaction_golden import load


def run(*texts):
    return compute([NodeIn(f"F1-P{i}", t) for i, t in enumerate(texts)], "purchase")


def test_01_sample():
    nodes, _, r, text, _ = load("names_derived_persons")
    for word in ("澄川虚构信息技术", "澄川虚构", "张经理", "李总", "Carter Fictive"):
        assert word not in text
    assert any(n.text[s.char_start:s.char_end] == "王建国" for s in r.suspects for n in nodes if n.loc == s.loc)


def test_02_short_derived():
    r = run("乙方：虚构科技有限公司", "虚构的方案")
    assert r.redacted["F1-P1"] == "虚构的方案"
    assert any(s.loc == "F1-P1" and s.reason == "derived_name" for s in r.suspects)


def test_03_industry_chain():
    r = run("乙方：东方虚构国际贸易有限公司", "东方虚构的报价，东方红")
    assert r.redacted["F1-P1"] == "公司_A的报价，东方红"


def test_04_no_title_fragment():
    assert not run("副总经理负责").items


def test_05_compound_surname():
    assert run("请欧阳总监确认").redacted["F1-P0"] == "请人员_1确认"


def test_06_english_title():
    assert run("Dear Dr. Lena Fictive,").redacted["F1-P0"] == "Dear Dr. 人员_1,"


def test_07_place():
    r = run("与张家港港务局联系")
    assert not any(s.char_start == 1 and s.char_end == 4 for s in r.suspects)


def test_09_chat_existing():
    assert _redact("澄川虚构的报价是否偏高？", {"公司_A": "澄川虚构信息技术有限公司"}, []).text == "公司_A的报价是否偏高？"


@pytest.mark.usefixtures("env")
def test_08_only_full_names_learned(db_sessionmaker, alice, ready_task):
    user, _, _ = alice
    doc = Document()
    doc.add_paragraph("乙方：澄川虚构信息技术有限公司")
    doc.add_paragraph("澄川虚构信息技术标准与澄川虚构的报价")
    data = io.BytesIO()
    doc.save(data)
    task = ready_task(user, files=[("虚构合同.docx", data.getvalue())])
    with db_sessionmaker() as db:
        learn_confirmed_names(db, task, user.id)
        names = {decrypt_text(row.term_enc) for row in db.scalars(select(RedactionTerm))}
        assert names == {"澄川虚构信息技术有限公司"}


@pytest.mark.parametrize("text", ["请王建国处理", "通知王建国，", "抄送欧阳建国。"])
def test_new_narrative_candidates(text):
    r = run(text)
    if text == '通知王建国，':
        # V2Q-44：仅此参数允许真实模型完整替换，不放宽其他参数或样本。
        start = text.index('王建国')
        suspected = any(s.char_start <= start and s.char_end >= start + len('王建国') for s in r.suspects)
        if not suspected:
            assert any(i.detected_by == 'ner' and i.category == 'PERSON'
                       and i.char_start <= start and i.char_end >= start + len('王建国') for i in r.items)
            assert '王建国' not in r.redacted['F1-P0']
    else:
        assert r.suspects


def test_overlap_keeps_existing_reason():
    assert any(s.reason == "person_name" for s in run("与王建国沟通").suspects)
