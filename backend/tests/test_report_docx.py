"""T-19：报告结构、快照保留、还原、XML 文本与真实数字。"""
import io
import re
from datetime import datetime, timezone

import pytest
from docx import Document
from sqlalchemy import event, select, update

from app.db.models import DocNode, Finding, RequirementResult, ReviewRun, Task, TaskFile
from app.export.report_docx import DECLARATION, SANCTIONS_NOTE, build_local_report, build_run_report
from app.security.crypto import encrypt_text
from tests.review_support import env, ready_task  # noqa: F401
from tests.test_checks_service import ROOT


@pytest.fixture
def report_case(make_user, ready_task, db_sessionmaker):
    user = make_user("report-owner")
    task = ready_task(user)
    with db_sessionmaker() as db:
        run = ReviewRun(task_id=task.id, run_no=2, status="completed", stance="purchase", source="incoming",
                        main_doc_type="contract", model_name="mock-model", builtin_req_version="2026.10-1",
                        company_req_version=3, created_by=user.id, finished_at=datetime(2026, 10, 20, 0, 12, tzinfo=timezone.utc),
                        doc_summary="我方与客户_A订立合同", sanctions_level="attention", sanctions_reasons=["客户_A需核实"])
        db.add(run)
        db.flush()
        db.add_all([
            Finding(run_id=run.id, task_id=task.id, seq=2, title="虚构问题", category="commercial", severity="low",
                    locs=["F1-P2"], cite_status="verified", quote_snapshot_enc=encrypt_text("星河虚构进出口有限公司快照"),
                    problem="客户_A需要明确义务\x01", impact="影响我方付款", basis="通行商业惯例",
                    suggestion="请客户_A补充", suggested_text="客户_A应按期交货", handling_status="adopted", handling_note="虚构处理备注"),
            Finding(run_id=run.id, task_id=task.id, seq=1, title="缺失项", category="legal", severity="high",
                    locs=[], missing=True, cite_status="missing", problem="需补充约定", handling_status="pending"),
            RequirementResult(run_id=run.id, req_id="C-06", req_source="builtin", req_text="虚构要求快照",
                              verdict="issue", note="客户_A需补充"),
        ])
        db.get(Task, task.id).status = "reviewed"
        db.get(Task, task.id).latest_run_id = run.id
        db.commit()
    return user, task, run


def contents(data):
    doc = Document(io.BytesIO(data))
    text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    return doc, text


def test_report_sections_restore_snapshot_and_layout(report_case, db_sessionmaker):
    user, task, run = report_case
    with db_sessionmaker() as db:
        filename, data = build_run_report(db, db.get(ReviewRun, run.id), user)
    doc, text = contents(data)
    headings = [p.text for p in doc.paragraphs if p.text[:1] in "一二三四五六七"]
    assert headings == ["一 基本信息", "二 覆盖说明与限制事项", "三 结论摘要", "四 问题清单", "五 本地校验结果", "六 审核要求逐条结论", "七 固定声明"]
    assert doc.paragraphs[0].text == "商务文件审核报告"
    assert str(doc.paragraphs[0].runs[0].font.color.rgb) == "000000"
    assert not doc.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    for value in (DECLARATION, SANCTIONS_NOTE, "星河虚构进出口有限公司快照", "Fictional Metals Trading FZE需要明确义务",
                  "虚构要求快照", "已采纳", "虚构处理备注", "公司补充要求 第 3 版", "第 2 轮", "2026-10-20 08:12", "对我方的影响：", "依据：", "建议措辞："):
        assert value in text
    assert "客户_A" not in text and "\x01" not in text and "审核通过" not in text and "无风险" not in text
    assert filename.startswith("审核报告_销售合同_") and filename.endswith(".docx")
    section = doc.sections[0]
    assert abs(section.page_width.cm - 29.7) < .01 and abs(section.page_height.cm - 21) < .01
    assert doc.styles["Normal"].font.size.pt == 10.5
    issues = doc.tables[1]
    assert [c.text for c in issues.rows[0].cells] == ["序号", "等级", "类别", "位置", "原文", "问题说明", "修改建议", "引用状态", "处理状态与备注"]
    assert issues.cell(1, 0).text == "1" and issues.cell(1, 4).text == "—"
    assert issues.cell(2, 0).text == "2" and "快照" in issues.cell(2, 4).text


def test_report_does_not_select_node_original_and_survives_purge(report_case, db_sessionmaker):
    user, task, run = report_case
    statements = []
    with db_sessionmaker() as db:
        db.execute(update(DocNode).where(DocNode.task_id == task.id).values(text_enc=None))
        db.commit()
        engine = db.get_bind()
        def collect(conn, cursor, statement, parameters, context, executemany):
            statements.append(statement)
        event.listen(engine, "before_cursor_execute", collect)
        try:
            _, data = build_run_report(db, db.get(ReviewRun, run.id), user)
        finally:
            event.remove(engine, "before_cursor_execute", collect)
    _, text = contents(data)
    assert "星河虚构进出口有限公司快照" in text
    assert not any("doc_nodes.text_enc" in sql.lower() for sql in statements)


def test_local_report_and_sanitized_filename(make_user, ready_task, db_sessionmaker):
    owner = make_user("local-report")
    task = ready_task(owner, [("报价单.xlsx", (ROOT / "row_error.xlsx").read_bytes())], policy="local_only")
    with db_sessionmaker() as db:
        file = db.scalar(select(TaskFile).where(TaskFile.task_id == task.id))
        file.original_name = '虚构/报\\价:*?"<>|\x01.xlsx'
        db.flush()
        filename, data = build_local_report(db, task, owner)
    doc, text = contents(data)
    assert "未做 AI 审核" in text and "52,400.00" in text and DECLARATION in text
    assert [p.text for p in doc.paragraphs if p.text[:1] in "一二三四五六七"] == ["一 基本信息", "五 本地校验结果", "七 固定声明"]
    assert not re.search(r'[\\/:*?"<>|\x00-\x1f]', filename)
    assert not any(phrase in text for phrase in ("制裁结论", "审核通过", "无风险"))


def test_empty_local_report(make_user, ready_task, db_sessionmaker):
    owner = make_user("empty-report")
    task = ready_task(owner, [("合同.docx", (ROOT / "delivery_ok.docx").read_bytes())], policy="local_only")
    with db_sessionmaker() as db:
        _, data = build_local_report(db, task, owner)
    assert "本地校验未发现问题" in contents(data)[1]


def test_export_filename_uses_shanghai_day(monkeypatch):
    from types import SimpleNamespace
    from app.export import report_docx

    class Clock:
        @staticmethod
        def now(tz):
            return datetime(2026, 10, 20, 20, 30, tzinfo=timezone.utc).astimezone(tz)

    monkeypatch.setattr(report_docx, "datetime", Clock)
    assert report_docx._filename(SimpleNamespace(original_name="虚构报价.xlsx"), None) == "审核报告_虚构报价_20261021.docx"
