"""R4-03：设计 4.6 的十一项契约，全部使用虚构内容与模拟模型。"""
import io
import json
import logging
import uuid
from datetime import timedelta

import psycopg
import pytest
from alembic import command
from docx import Document
from sqlalchemy import select

from app.admin import settings_store
from app.backup import retention
from app.db.base import utcnow
from app.db.models import DocNode, Job, LlmCall, RedactionItem, RedactionSuspect, RequirementResult, ReviewRun, Task
from app.main import API_PREFIX
from app.redaction import engine
from app.review import job, service
from app.review.prompt import build_messages
from app.security import crypto
from tests.review_support import env, llm, ready_task  # noqa: F401
from tests.test_migrations import ADMIN_URL, _alembic, _insert_task, _insert_user, _url_for

pytestmark = pytest.mark.usefixtures("env")
NOTE = "这是给南海虚构贸易有限公司的报价，联系人：周示例，合同总价：98000"
SECRETS = ("南海虚构贸易有限公司", "周示例", "98000")
SUSPECT = "Meridian Fictive Inspection"


@pytest.fixture
def alice(make_user, login):
    user = make_user("alice")
    client, csrf = login("alice")
    return user, client, {"X-CSRF-Token": csrf}


def task_url(task, suffix=""):
    return f"{API_PREFIX}/tasks/{task.id}{suffix}"


def complete(sm, client, headers, task, llm):
    requests = llm("normal")
    response = client.post(task_url(task, "/runs"), headers=headers)
    assert response.status_code == 201, response.text
    rid = uuid.UUID(response.json()["id"])
    with sm() as db:
        job.run_review(db, rid)
        db.commit()
        assert db.get(ReviewRun, rid).status in ("completed", "partial")
    return rid, requests


def test_r4_note_outbound_and_storage(db_sessionmaker, alice, ready_task, llm):
    """4.6 #1、#2：NOTE 与外发记录均为代号；逐表逐列检查数据库无三处原文。"""
    user, c, h = alice
    task = ready_task(user, user_note=NOTE)
    preview = c.get(task_url(task, "/redaction/outbound")).json()
    assert preview["note"] and "[NOTE]" not in preview["text"]
    assert all(secret not in preview["note"] for secret in SECRETS)
    assert {i["category"] for i in c.get(task_url(task, "/redaction")).json()["items"] if i["loc"] == "NOTE"} == {"ORG", "PERSON", "AMOUNT"}
    assert c.get(task_url(task)).json()["user_note"] == NOTE
    assert any(p["name"] == SECRETS[0] and p["role_label"] == "其他" for p in c.get(task_url(task, "/parties")).json()["items"])
    rid, requests = complete(db_sessionmaker, c, h, task, llm)
    assert len(requests) == 1
    wire = json.loads(requests[0].content)
    note_lines = [line for m in wire["messages"] for line in m["content"].splitlines() if "[NOTE]" in line]
    assert len(note_lines) == 1 and preview["note"] in note_lines[0]
    with db_sessionmaker() as db:
        for call in db.scalars(select(LlmCall)):
            assert all(secret not in call.request_text for secret in SECRETS)
        for model in (Task, ReviewRun, RequirementResult):
            rows = db.execute(select(model.__table__)).mappings().all()
            assert rows
            for row in rows:
                for column, value in row.items():
                    for secret in SECRETS:
                        assert secret not in str(value), (model.__tablename__, column)
                        if isinstance(value, bytes):
                            assert secret.encode() not in value
        stored = db.get(Task, task.id)
        assert crypto.decrypt_text(stored.user_note_enc) == NOTE
        assert db.get(ReviewRun, rid).user_note == preview["note"]


def test_r4_note_suspect_blocks_send(alice, ready_task, db_sessionmaker):
    """4.6 #3：NOTE 疑似遗漏及上下文，未处理时不创建审核或模型记录。"""
    user, c, h = alice
    task = ready_task(user, user_note=SUSPECT)
    found = c.get(task_url(task, "/redaction")).json()["suspects"]
    assert [(s["loc"], s["text"], s["decision"]) for s in found] == [("NOTE", SUSPECT, "pending")]
    response = c.post(task_url(task, "/runs"), headers=h)
    assert (response.status_code, response.json()["code"]) == (409, "SUSPECTS_PENDING")
    with db_sessionmaker() as db:
        assert list(db.scalars(select(ReviewRun))) == []
        assert list(db.scalars(select(LlmCall))) == []


def test_r4_note_changed_resets_old_decisions(alice, ready_task, db_sessionmaker):
    """4.6 #4：等长新文字不能继承原位置的忽略决定，未改变文字时决定保留。"""
    user, c, h = alice
    task = ready_task(user, user_note=SUSPECT)
    sid = c.get(task_url(task, "/redaction")).json()["suspects"][0]["id"]
    assert c.post(task_url(task, f"/redaction/suspects/{sid}"), json={"decision": "ignore"}, headers=h).status_code == 200
    assert c.patch(task_url(task), json={"user_note": SUSPECT}, headers=h).status_code == 200
    assert c.get(task_url(task, "/redaction")).json()["suspects"][0]["decision"] == "ignore"
    replacement = "Northern Fictive Inspection"
    assert len(replacement) == len(SUSPECT)
    assert c.patch(task_url(task), json={"user_note": replacement}, headers=h).status_code == 200
    found = c.get(task_url(task, "/redaction")).json()["suspects"]
    assert [(s["text"], s["decision"]) for s in found] == [(replacement, "pending")]
    with db_sessionmaker() as db:
        assert not db.scalar(select(RedactionSuspect).where(RedactionSuspect.task_id == task.id, RedactionSuspect.id == sid))


def test_r4_note_manual_term_is_shared_and_cleared(alice, ready_task, db_sessionmaker):
    """4.6 #5：NOTE 手动词条全文生效；说明改动后删除旧位置词条。"""
    user, c, h = alice
    doc = Document()
    doc.add_paragraph("卖方：星河虚构进出口有限公司")
    doc.add_paragraph("品名：针状石油焦")
    data = io.BytesIO()
    doc.save(data)
    task = ready_task(user, files=[("虚构合同.docx", data.getvalue())], user_note="请核对针状石油焦")
    response = c.post(task_url(task, "/redaction/manual"), headers=h,
                      json={"loc": "NOTE", "char_start": 3, "char_end": 8, "category": "CUSTOM"})
    assert response.status_code == 200, response.text
    preview = c.get(task_url(task, "/redaction/outbound")).json()
    assert "针状石油焦" not in preview["text"] and "针状石油焦" not in preview["note"]
    with db_sessionmaker() as db:
        items = list(db.scalars(select(RedactionItem).where(RedactionItem.task_id == task.id, RedactionItem.detected_by == "user")))
        assert any(i.loc == "NOTE" for i in items) and any(i.loc != "NOTE" for i in items)
        assert len({i.placeholder for i in items}) == 1
    assert c.patch(task_url(task), json={"user_note": "请检查交付条款"}, headers=h).status_code == 200
    with db_sessionmaker() as db:
        assert not db.scalar(select(RedactionItem).where(RedactionItem.task_id == task.id, RedactionItem.loc == "NOTE", RedactionItem.detected_by == "user"))


def test_r4_note_result_report_and_retention(db_sessionmaker, alice, ready_task, llm):
    """4.6 #6、#9：结果与报告还原，清理说明密文后结果和报告仍可还原。"""
    user, c, h = alice
    task = ready_task(user, user_note=NOTE)
    rid, _requests = complete(db_sessionmaker, c, h, task, llm)
    for after_retention in (False, True):
        if after_retention:
            with db_sessionmaker() as db:
                settings_store.set_setting(db, "retention.original_days", 1, None)
                db.get(Task, task.id).last_completed_at = utcnow() - timedelta(days=2)
                db.commit()
                assert retention.run_retention(db) == {"tasks": 1, "files": 1}
                stored = db.get(Task, task.id)
                assert stored.user_note_enc is None and stored.user_note_redacted
        detail = c.get(f"{API_PREFIX}/runs/{rid}").json()
        assert detail["user_note"] == NOTE
        assert next(r["req_text"] for r in detail["requirement_results"] if r["req_id"] == "NOTE") == NOTE
        report = c.get(f"{API_PREFIX}/runs/{rid}/report")
        assert report.status_code == 200
        doc = Document(io.BytesIO(report.content))
        assert NOTE in "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)


def test_r4_empty_note_keeps_original_messages(db_sessionmaker, alice, ready_task, llm):
    """4.6 #7：空说明沿用未修改的 prompt 构建结果，正文及节点不出现 NOTE。"""
    user, c, h = alice
    task = ready_task(user)
    with db_sessionmaker() as db:
        current = db.get(Task, task.id)
        assert current.user_note_enc is None and current.user_note_redacted is None
        ctx = service.build_context(db, current, stance=task.stance, source=task.source, user_note=None, main_doc_type="contract")
        expected = build_messages(ctx)
        assert all(n.loc != "NOTE" for n in db.scalars(select(DocNode)))
    assert c.get(task_url(task, "/redaction/outbound")).json()["note"] is None
    _rid, requests = complete(db_sessionmaker, c, h, task, llm)
    assert json.loads(requests[0].content)["messages"] == expected
    assert all("[NOTE]" not in m["content"] for m in expected)


def test_r4_note_migration_upgrade_downgrade():
    """4.6 #8：独立临时库，在新测试中覆盖数据处理，不修改既有迁移测试逻辑。"""
    name = "aireview_test_" + uuid.uuid4().hex[:10]
    with psycopg.connect(ADMIN_URL, autocommit=True) as admin:
        admin.execute(f'CREATE DATABASE "{name}"')
    try:
        url = _url_for(name)
        cfg = _alembic(url)
        command.upgrade(cfg, "0004")
        with psycopg.connect(url, autocommit=True) as conn:
            owner = _insert_user(conn)
            tid = _insert_task(conn, owner)
            conn.execute("UPDATE tasks SET user_note = %s WHERE id = %s", (NOTE, tid))
            rid = conn.execute("INSERT INTO review_runs (task_id, run_no, status, stance, source, main_doc_type, created_by, queued_at, user_note) VALUES (%s, 1, 'queued', 'sales', 'outgoing', 'contract', %s, now(), %s) RETURNING id", (tid, owner, NOTE)).fetchone()[0]
            conn.execute("INSERT INTO requirement_results (run_id, req_id, req_source, req_text, verdict) VALUES (%s, 'NOTE', 'note', %s, 'ok')", (rid, NOTE))
        command.upgrade(cfg, "0005")
        with psycopg.connect(url, autocommit=True) as conn:
            assert conn.execute("SELECT user_note_enc, user_note_redacted FROM tasks").fetchone() == (None, None)
            assert conn.execute("SELECT user_note FROM review_runs").fetchone() == (None,)
            assert conn.execute("SELECT req_text FROM requirement_results").fetchone() == ("（本次补充说明，升级时已清除）",)
            conn.execute("UPDATE review_runs SET user_note = %s", ("字" * 550,))
        command.downgrade(cfg, "0004")
        with psycopg.connect(url, autocommit=True) as conn:
            assert conn.execute("SELECT user_note FROM tasks").fetchone() == (None,)
            assert conn.execute("SELECT user_note FROM review_runs").fetchone() == ("字" * 500,)
            columns = {r[0] for r in conn.execute("SELECT column_name FROM information_schema.columns WHERE table_name='tasks'")}
            assert not {"user_note_enc", "user_note_redacted"} & columns
        command.upgrade(cfg, "head")
    finally:
        with psycopg.connect(ADMIN_URL, autocommit=True) as admin:
            admin.execute(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)')


@pytest.mark.parametrize("role", ["user", "admin"])
def test_r4_note_owner_isolation(db_sessionmaker, alice, ready_task, llm, make_user, login, role):
    """4.6 #10：新字段、NOTE 手动入口、轮次与报告均沿用归属 404。"""
    user, c, h = alice
    task = ready_task(user, user_note=NOTE)
    rid, _requests = complete(db_sessionmaker, c, h, task, llm)
    make_user("other", role=role)
    other, csrf = login("other")
    headers = {"X-CSRF-Token": csrf}
    for suffix in ("", "/redaction", "/redaction/outbound"):
        assert other.get(task_url(task, suffix)).status_code == 404
    assert other.patch(task_url(task), json={"user_note": "说明"}, headers=headers).status_code == 404
    assert other.post(task_url(task, "/redaction/manual"), headers=headers,
                      json={"loc": "NOTE", "char_start": 0, "char_end": 2, "category": "CUSTOM"}).status_code == 404
    for suffix in ("", "/report"):
        assert other.get(f"{API_PREFIX}/runs/{rid}{suffix}").status_code == 404


def test_r4_note_log_safety(db_sessionmaker, alice, ready_task, llm, caplog):
    """4.6 #11：DEBUG 捕获解析、确认、审核全流程。"""
    caplog.set_level(logging.DEBUG)
    user, c, h = alice
    task = ready_task(user, user_note=NOTE)
    complete(db_sessionmaker, c, h, task, llm)
    logs = "\n".join(r.getMessage() + (r.exc_text or "") for r in caplog.records)
    assert NOTE not in logs
    assert all(s not in logs for s in SECRETS)


def test_r4_note_keeps_document_codes_and_clears_note(db_sessionmaker, alice, ready_task):
    """NOTE 末尾分配代号，清空说明时不残留伪节点或原文。"""
    user, c, h = alice
    baseline = ready_task(user)
    task = ready_task(user, user_note=NOTE)
    with db_sessionmaker() as db:
        before = [(n.loc, n.redacted_text) for n in db.scalars(select(DocNode).where(DocNode.task_id == baseline.id).order_by(DocNode.loc))]
        after = [(n.loc, n.redacted_text) for n in db.scalars(select(DocNode).where(DocNode.task_id == task.id).order_by(DocNode.loc))]
        assert before == after
    assert c.patch(task_url(task), json={"user_note": ""}, headers=h).status_code == 200
    with db_sessionmaker() as db:
        stored = db.get(Task, task.id)
        assert stored.user_note_enc is None and stored.user_note_redacted is None
        assert list(db.scalars(select(RedactionItem).where(RedactionItem.task_id == task.id, RedactionItem.loc == "NOTE"))) == []
        assert list(db.scalars(select(RedactionSuspect).where(RedactionSuspect.task_id == task.id, RedactionSuspect.loc == "NOTE"))) == []


def test_r4_note_amount_is_not_contract_total():
    """R4 审核补充（4.1）：补充说明中的金额不作为合同总价、自身不加比例，正文金额不因此出现比例。"""
    body = [engine.NodeIn("F1-P1", "卖方：澄川虚构信息技术有限公司"), engine.NodeIn("F1-P2", "首期付款人民币30,000元。")]
    plain = engine.compute(body, "sales", main_prefix="F1-")
    with_note = engine.compute([*body, engine.NodeIn("NOTE", "合同总价：人民币100,000元")], "sales", main_prefix="F1-")
    assert with_note.redacted["F1-P2"] == plain.redacted["F1-P2"] == "首期付款人民币[金额_1]元。"
    assert "合同总价的" not in with_note.redacted["NOTE"]
