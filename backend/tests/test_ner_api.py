"""B-06：真实解析接口、追问确认链与模型来源不可撤销；云模型只用模拟网关。"""
import io
import json

import pytest
from docx import Document
from sqlalchemy import select

from app.db.models import LlmCall, RedactionItem
from app.redaction import ner
from tests.review_support import env, llm, ready_task  # noqa: F401
from tests.test_user_note_redaction import alice, complete, task_url  # noqa: F401
from tests.acceptance_support import settle_suspects

pytestmark = pytest.mark.usefixtures('env')


def test_real_model_body_note_chat_preview_and_send(alice, ready_task, db_sessionmaker, llm):
    user, client, headers = alice
    doc = Document()
    doc.add_paragraph('卖方：星河虚构进出口有限公司')
    doc.add_paragraph('沈柏宁复核了江岚研究院的检测计划。')
    data = io.BytesIO()
    doc.save(data)
    task = ready_task(user, files=[('虚构说明.docx', data.getvalue())],
                      user_note='陆清禾复核了澄岳实验室的意见。')
    response = client.get(task_url(task, '/redaction'))
    assert response.status_code == 200
    with db_sessionmaker() as db:
        model_items = list(db.scalars(select(RedactionItem).where(RedactionItem.task_id == task.id,
                                                                 RedactionItem.detected_by == 'ner')))
        assert model_items and any(i.loc == 'NOTE' for i in model_items)
        target = {k: getattr(model_items[0], k) for k in ('loc', 'char_start', 'char_end')}
    denied = client.post(task_url(task, '/redaction/revoke'), headers=headers,
                         json={k: target[k] for k in ('loc', 'char_start', 'char_end')})
    assert (denied.status_code, denied.json()['code']) == (422, 'NOT_REVOCABLE')
    settle_suspects(client, headers, task.id, [])
    complete(db_sessionmaker, client, headers, task, llm)
    message = '经与王建国沟通……'
    preview = client.post(task_url(task, '/chat/preview'), headers=headers, json={'message': message})
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert '王建国' not in body['redacted'] or any(s['text'] == '王建国' for s in body['suspects'])
    decisions = [{'start': s['start'], 'end': s['end'], 'decision': 'replace', 'category': 'PERSON'}
                 for s in body['suspects']]
    requests = llm('normal')
    sent = client.post(task_url(task, '/chat'), headers=headers,
                       json={'message': message, 'decisions': decisions})
    assert sent.status_code == 200, sent.text
    assert requests and '王建国' not in json.dumps(json.loads(requests[0].content), ensure_ascii=False)
    with db_sessionmaker() as db:
        calls = list(db.scalars(select(LlmCall).where(LlmCall.task_id == task.id)))
        assert all('王建国' not in call.request_text for call in calls)
        assert db.scalar(select(RedactionItem.id).where(RedactionItem.task_id == task.id,
                                                       RedactionItem.detected_by == 'ner')) is not None


def test_medium_model_decision_recomputed_at_send(alice, ready_task, db_sessionmaker, llm, monkeypatch):
    user, client, headers = alice
    task = ready_task(user)
    settle_suspects(client, headers, task.id, [])
    complete(db_sessionmaker, client, headers, task, llm)
    message = '梓澄复核。'

    def model(text):
        return [ner.Entity(0, 2, 'PERSON', 0.7)] if text == message else []

    monkeypatch.setattr(ner, 'recognize', model)
    preview = client.post(task_url(task, '/chat/preview'), headers=headers, json={'message': message})
    assert preview.json()['suspects'] == [{'start': 0, 'end': 2, 'text': '梓澄', 'reason': 'ner'}]
    requests = llm('normal')
    blocked = client.post(task_url(task, '/chat'), headers=headers, json={'message': message})
    assert (blocked.status_code, blocked.json()['code']) == (409, 'CHAT_SUSPECTS_PENDING')
    assert not requests
    sent = client.post(task_url(task, '/chat'), headers=headers, json={'message': message,
        'decisions': [{'start': 0, 'end': 2, 'decision': 'replace', 'category': 'PERSON'}]})
    assert sent.status_code == 200, sent.text
    assert '梓澄' not in json.dumps(json.loads(requests[0].content), ensure_ascii=False)
