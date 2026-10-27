"""T-44 2、3：虚构法规逐条拆分。"""
import io

from docx import Document

from app.laws.articles import split
from app.parsing.docx import parse_docx


def law_doc():
    doc = Document()
    for text in ['示例供货合同法', '（示例题注）', '目录', '第一编　总则', '第一章　范围',
                 '第一条　示例适用范围。', '第二款独立段落。', '第三十条规定的内容继续适用。',
                 '第二章　约定', '第一百二十条之一　示例交付。', '一、示例解释分部',
                 '第一千零一条　示例责任。', '附 则']:
        doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '示例附表文字。'
    doc.sections[0].header.paragraphs[0].text = '第二条　页眉不参与。'
    doc.sections[0].footer.paragraphs[0].text = '第三条　页脚不参与。'
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def test_t44_2_3_articles_exact():
    got = split(parse_docx(law_doc(), 1).nodes)
    assert [(a.article_no, a.seq, a.text) for a in got] == [
        ('第一条', 1, '示例适用范围。\n第二款独立段落。\n第三十条规定的内容继续适用。'),
        ('第一百二十条之一', 2, '示例交付。'),
        ('第一千零一条', 3, '示例责任。\n示例附表文字。')]


def test_no_articles():
    got = split(parse_docx(_doc_without_articles(), 1).nodes, title='示例法')
    assert [(a.article_no, a.seq, a.text) for a in got] == [(None, 1, '示例前言')]


def _doc_without_articles():
    d = Document()
    d.add_paragraph('示例前言')
    o = io.BytesIO()
    d.save(o)
    return o.getvalue()


def unnumbered_doc(lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.sections[0].header.paragraphs[0].text = '页眉不得参与。'
    doc.sections[0].footer.paragraphs[0].text = '页脚不得参与。'
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def segments(lines, title='示例批复'):
    return split(parse_docx(unnumbered_doc(lines), 1).nodes, title=title)


def test_t44a_1_reply_preamble_and_trailing_date():
    body = ['示例省高级人民法院：经研究，批复如下：', '应依法处理示例事项。', '此复。', '二〇二〇年三月十九日']
    got = segments(['示例批复', '（示例题注）', '法释〔2020〕99号', '二〇二〇年三月十八日', '示例公告'] + body)
    assert [(a.article_no, a.seq, a.text) for a in got] == [(None, 1, '\n'.join(body))]


def test_t44a_2_chinese_items():
    lines = ['一、示例范围。', '二、示例程序。', '三、示例责任。']
    assert [(a.article_no, a.seq, a.text) for a in segments(lines)] == [
        (None, i, line) for i, line in enumerate(lines, 1)]


def test_t44a_3_merge_headings_and_preserve_numbers():
    got = segments(['一、执行机构', '1．示例职责。', '２．示例程序。', '二、执行管辖', '3.示例管辖。'])
    assert [a.text for a in got] == ['一、执行机构\n1．示例职责。', '２．示例程序。', '二、执行管辖\n3.示例管辖。']
    assert all(a.article_no is None for a in got)
    assert [a.text for a in segments(['正文内容。', '一、末尾标题'])] == ['正文内容。\n一、末尾标题']


def test_t44a_4_paragraph_boundary_limit():
    lines = [str(i) * 199 + '。' for i in range(5)]
    got = segments(lines + ['长' * 899 + '。'])
    assert [a.text for a in got] == ['\n'.join(lines[:3]), '\n'.join(lines[3:]), '长' * 899 + '。']
    assert [a.seq for a in got] == [1, 2, 3]


def test_t44a_5_no_space_contiguous_articles():
    got = segments(['第一条依法处理。', '第一条之一补充规定。', '第二条依约处理。',
                    '第五条规定的内容继续适用。', '第三条示例责任。'])
    assert [(a.article_no, a.text) for a in got] == [('第一条', '依法处理。'),
        ('第一条之一', '补充规定。'), ('第二条', '依约处理。\n第五条规定的内容继续适用。'), ('第三条', '示例责任。')]
    assert [(a.article_no, a.text) for a in segments(['第一条之一不得作为首条。'])] == [
        (None, '第一条之一不得作为首条。')]


def test_t44a_6_normal_articles_keep_strict_matching():
    got = segments(['第一条　正常正文。', '第三十条规定的继续适用。', '第二条无空白也不得单独识别。'])
    assert [(a.article_no, a.text) for a in got] == [('第一条', '正常正文。\n第三十条规定的继续适用。\n第二条无空白也不得单独识别。')]


def test_t44a_7_empty_after_preamble():
    assert segments(['示例批复', '（示例题注）', '法释〔2020〕99号', '二〇二〇年三月十九日']) == []


def test_t44a_12_no_content_in_logs(caplog):
    segments(['示例批复', '一、仅供日志检查的示例正文。'])
    assert '示例批复' not in caplog.text and '仅供日志检查' not in caplog.text
