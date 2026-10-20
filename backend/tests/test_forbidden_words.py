"""A-08：扫描系统自有文字，避免界面和报告出现禁用措辞。"""
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import pytest


ROOT = Path(__file__).resolve().parents[2]
FORBIDDEN_WORDS = ("审核通过", "无风险")
SANCTIONS_LABEL = "无明显风险"
REPORT_TEMPLATE_DIR_NAMES = {"report_template", "report_templates", "报告模板"}


def _is_report_template_path(path: Path, backend_root: Path) -> bool:
    parts = path.relative_to(backend_root).parts[:-1]
    normalized = [part.casefold().replace("-", "_") for part in parts]
    if any(part in REPORT_TEMPLATE_DIR_NAMES for part in normalized):
        return True
    return any(
        part == "templates" and index > 0 and normalized[index - 1] in {"report", "reports"}
        for index, part in enumerate(normalized)
    )


def _is_test_or_fixture(path: Path, root: Path) -> bool:
    relative = path.relative_to(root)
    excluded_dirs = {"test", "tests", "__tests__", "fixture", "fixtures"}
    if any(
        part.casefold() in excluded_dirs for part in relative.parts[:-1]
    ):
        return True

    filename = path.name.casefold()
    stem = path.stem.casefold()
    return (
        filename.startswith("test_")
        or stem.endswith("_test")
        or stem.endswith(".test")
        or stem.endswith(".spec")
    )


def scan_files(root: Path) -> list[Path]:
    """返回设计第十节规定的源码、资源与可选报告模板文件。"""
    frontend_src = root / "frontend" / "src"
    backend_root = root / "backend"
    app_dir = backend_root / "app"
    resources_dir = backend_root / "resources"

    files = {
        path
        for path in frontend_src.rglob("*")
        if path.is_file() and path.suffix.casefold() in {".vue", ".ts"}
    }
    files.update(path for path in app_dir.rglob("*.py") if path.is_file())
    files.update(path for path in resources_dir.rglob("*") if path.is_file())
    files.update(
        path
        for path in backend_root.rglob("*")
        if path.is_file() and _is_report_template_path(path, backend_root)
    )
    return sorted(path for path in files if not _is_test_or_fixture(path, root))


def _read_text(path: Path) -> str:
    if path.suffix.casefold() != ".docx":
        return path.read_text(encoding="utf-8")

    # 报告模板可能是 Word 文件；读取其 XML 文本，避免把压缩后的二进制直接当文本扫描。
    text_parts = []
    word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                document = ET.fromstring(archive.read(name))
                for paragraph in document.iter(f"{word_ns}p"):
                    text_parts.append(
                        "".join(node.text or "" for node in paragraph.iter(f"{word_ns}t"))
                    )
    return "\n".join(text_parts)


def _forbidden_words(text: str) -> list[str]:
    # “无明显风险”是 SPEC 11.3 的制裁等级，按设计约定先移除再检查。
    normalized = text.replace(SANCTIONS_LABEL, "")
    return [word for word in FORBIDDEN_WORDS if word in normalized]


def test_scan_scope_includes_sources_and_excludes_tests_and_fixtures(tmp_path):
    expected = {
        tmp_path / "frontend" / "src" / "App.vue",
        tmp_path / "frontend" / "src" / "main.ts",
        tmp_path / "backend" / "app" / "main.py",
        tmp_path / "backend" / "resources" / "prompts" / "system.md",
        tmp_path / "backend" / "report_templates" / "report.txt",
    }
    for path in expected:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("示例内容", encoding="utf-8")

    ignored = (
        tmp_path / "frontend" / "src" / "ignored.js",
        tmp_path / "frontend" / "src" / "sample.test.ts",
        tmp_path / "backend" / "app" / "ignored.md",
        tmp_path / "backend" / "app" / "test_helper.py",
        tmp_path / "backend" / "tests" / "test_sample.py",
        tmp_path / "backend" / "tests" / "fixtures" / "sample.txt",
        tmp_path / "backend" / "resources" / "fixtures" / "sample.txt",
        tmp_path / "backend" / "tests" / "report_templates" / "report.txt",
    )
    for path in ignored:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("示例内容", encoding="utf-8")

    assert set(scan_files(tmp_path)) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("该事项审核通过", ["审核通过"]),
        ("这是无风险结论", ["无风险"]),
        ("制裁等级：无明显风险", []),
        ("合规审查完成", []),
    ],
)
def test_forbidden_word_detector(text, expected):
    assert _forbidden_words(text) == expected


def test_word_report_template_text_is_scanned(tmp_path):
    from docx import Document

    template = tmp_path / "backend" / "report_templates" / "template.docx"
    template.parent.mkdir(parents=True)
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("审核")
    paragraph.add_run("通过")
    document.save(template)

    assert template in scan_files(tmp_path)
    assert _forbidden_words(_read_text(template)) == ["审核通过"]


def test_project_system_text_does_not_contain_forbidden_words():
    # Docker 执行时需把 frontend/ 只读挂载到 /srv/frontend，与 backend/ 同级。
    files = scan_files(ROOT)
    assert ROOT / "frontend" / "src" / "App.vue" in files
    assert ROOT / "backend" / "app" / "main.py" in files
    assert ROOT / "backend" / "resources" / "prompts" / "system.md" in files

    violations = {}
    for path in files:
        hits = _forbidden_words(_read_text(path))
        if hits:
            violations[str(path.relative_to(ROOT))] = hits
    assert violations == {}, f"系统文字含禁用词：{violations}"
