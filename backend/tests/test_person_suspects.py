"""R4-08：设计 11.6 逐条契约；仅用虚构姓名和模拟网关。"""
import io
import json
import logging
import uuid

import psycopg
import pytest
from alembic import command
from docx import Document

from app.main import API_PREFIX
from app.redaction.engine import load_redaction_map
from app.redaction.recheck import recheck
from app.redaction.suspects import find_suspects
from tests.review_support import env, llm, ready_task  # noqa: F401
from tests.test_chat import alice, reviewed, chat_reply, chat_calls, code_of  # noqa: F401
from tests.test_migrations import ADMIN_URL, _alembic, _insert_task, _insert_user, _url_for
from tests.test_user_note_redaction import complete, task_url


@pytest.mark.parametrize(('text', 'name'), [
    ('对方联系人为陆虚岚，负责验收。', '陆虚岚'),
    ('陆虚岚提到的预付款：45000 是否偏高？', '陆虚岚'),
    ('陆虚岚说预付款太高。', '陆虚岚'),
    ('请与欧阳虚明联系确认交期。', '欧阳虚明'),
    ('张虚强经理要求我们尽快回复。', '张虚强'),
    ('对接人是何虚芳，她负责验收。', '何虚芳'),
    ('昨天王虚一表示可以接受。', '王虚一'),
    ('对方王虚一表示同意。', '王虚一'),
    ('经办人为林虚，请核对。', '林虚'),
    ('宋虚婷女士已确认样品合格。', '宋虚婷'),
])
def test_design_11_6_positive(text, name):
    assert [(text[s.start:s.end], s.reason) for s in find_suspects(text, [])] == [(name, 'person_name')]


@pytest.mark.parametrize('text', [
    '王总提到的条款', '李经理和赵总监都同意了', '周期表示为 30 天', '常规要求见第三条',
    '我们周五前回复', '本合同项目名称为虚构工程', '方案说明见附件', '文件说明见附件二',
    '双方同意按本合同执行', '任何一方提出异议', '由于对方要求提前交货', '黄金价格表示上涨',
    '高于市场价格的部分由买方承担', '单位说明保留', '公司要求付款在 30 天内', '王某表示同意',
])
def test_design_11_6_negative(text):
    assert not [s for s in find_suspects(text, []) if s.reason == 'person_name']


def test_design_11_6_occupied():
    text = '联系人：周示例，周示例经理同意。'
    assert find_suspects(text, [(4, 7), (8, 11)]) == []


@pytest.mark.usefixtures('env')
@pytest.mark.parametrize('in_note', [False, True])
def test_design_11_6_body_note_gateway_logs(db_sessionmaker, alice, ready_task, llm, caplog, in_note):
    """#13、14、18：真实解析、接口、阻断、替换、模拟网关、出口复检与 DEBUG 日志。"""
    caplog.set_level(logging.DEBUG)
    user, c, h = alice
    sentence = '对方联系人为陆虚岚' if in_note else '对方联系人为陆虚岚，负责验收。'
    doc = Document()
    doc.add_paragraph('卖方：星河虚构进出口有限公司')
    if not in_note:
        doc.add_paragraph(sentence)
    data = io.BytesIO()
    doc.save(data)
    task = ready_task(user, files=[('虚构备忘.docx', data.getvalue())], user_note=sentence if in_note else None)
    suspects = c.get(task_url(task, '/redaction')).json()['suspects']
    assert len(suspects) == 1
    s = suspects[0]
    assert (s['loc'], s['text'], s['reason']) == ('NOTE' if in_note else 'F1-P2', '陆虚岚', 'person_name')
    blocked = c.post(task_url(task, '/runs'), headers=h)
    assert (blocked.status_code, blocked.json()['code']) == (409, 'SUSPECTS_PENDING')
    r = c.post(task_url(task, f"/redaction/suspects/{s['id']}"), headers=h,
               json={'decision': 'replace', 'category': 'PERSON'})
    assert r.status_code == 200, r.text
    outbound = c.get(task_url(task, '/redaction/outbound')).json()
    value = outbound['note' if in_note else 'text']
    with db_sessionmaker() as db:
        mapping = load_redaction_map(db, task.id)
    code = next(k for k, v in mapping.items() if v == '陆虚岚')
    assert code.startswith('人员_') and code in value and '陆虚岚' not in value
    assert recheck(value, real_values=mapping.values(), terms=[]) == []
    _, requests = complete(db_sessionmaker, c, h, task, llm)
    assert requests and all('陆虚岚' not in req.content.decode() for req in requests)
    wire = json.loads(requests[0].content)
    assert any(code in m['content'] for m in wire['messages'])
    assert '陆虚岚' not in caplog.text


@pytest.mark.usefixtures('env')
@pytest.mark.parametrize('decision', ['replace', 'keep'])
def test_design_11_6_chat_decisions(db_sessionmaker, alice, ready_task, llm, caplog, decision):
    """#15、18：无决定不进网关，替换及保留分别在独立任务验证。"""
    caplog.set_level(logging.DEBUG)
    user, c, h = alice
    task = reviewed(db_sessionmaker, user, c, h, ready_task, llm)
    url = task_url(task, '/chat')
    message = '陆虚岚提到的预付款是否偏高？'
    p = c.post(url + '/preview', json={'message': message}, headers=h)
    assert p.status_code == 200
    assert p.json()['suspects'] == [{'start': 0, 'end': 3, 'text': '陆虚岚', 'reason': 'person_name'}]
    sent = []
    llm(handler=chat_reply('[人员_Q1]的建议应复核。' if decision == 'replace' else '建议复核。', sent))
    blocked = c.post(url, json={'message': message}, headers=h)
    assert (blocked.status_code, blocked.json()['code']) == (409, 'CHAT_SUSPECTS_PENDING')
    assert sent == [] and chat_calls(db_sessionmaker) == []
    d = {'start': 0, 'end': 3, 'decision': decision}
    if decision == 'replace':
        d['category'] = 'PERSON'
    reply = c.post(url, json={'message': message, 'decisions': [d]}, headers=h)
    assert reply.status_code == 200, reply.text
    wire = sent[0]['messages'][-1]['content']
    if decision == 'replace':
        assert wire == message.replace('陆虚岚', '[人员_Q1]')
        assert reply.json()['assistant']['content'] == '陆虚岚的建议应复核。'
    else:
        assert wire == message
    assert '陆虚岚' not in caplog.text


@pytest.mark.usefixtures('env')
def test_design_11_6_existing_person_map(db_sessionmaker, alice, ready_task, llm):
    user, c, h = alice
    doc = Document()
    doc.add_paragraph('卖方：星河虚构进出口有限公司')
    doc.add_paragraph('联系人：周示例')
    data = io.BytesIO()
    doc.save(data)
    task = reviewed(db_sessionmaker, user, c, h, ready_task, llm, files=[('虚构合同.docx', data.getvalue())])
    p = c.post(task_url(task, '/chat/preview'), json={'message': '周示例提到的条款'}, headers=h)
    assert p.status_code == 200
    assert p.json() == {'redacted': code_of(db_sessionmaker, task.id, '周示例') + '提到的条款', 'changed': True, 'suspects': []}


@pytest.mark.skipif(not ADMIN_URL, reason='未设置 TEST_DATABASE_URL')
def test_design_11_6_migration():
    """#17：真实约束写入、降级删新原因并保留旧原因、恢复旧约束。"""
    name = 'aireview_test_' + uuid.uuid4().hex[:10]
    with psycopg.connect(ADMIN_URL, autocommit=True) as admin:
        admin.execute(f'CREATE DATABASE "{name}"')
    try:
        url = _url_for(name)
        cfg = _alembic(url)
        command.upgrade(cfg, '0005')
        command.upgrade(cfg, '0006')
        with psycopg.connect(url, autocommit=True) as conn:
            tid = _insert_task(conn, _insert_user(conn))
            sql = "INSERT INTO redaction_suspects (task_id, loc, char_start, char_end, reason) VALUES (%s, 'F1-P1', 0, 3, %s)"
            conn.execute(sql, (tid, 'person_name'))
            conn.execute(sql, (tid, 'capitalized'))
        command.downgrade(cfg, '0005')
        with psycopg.connect(url, autocommit=True) as conn:
            assert conn.execute('SELECT reason FROM redaction_suspects').fetchall() == [('capitalized',)]
            with pytest.raises(psycopg.errors.CheckViolation):
                conn.execute(sql, (tid, 'person_name'))
        command.upgrade(cfg, 'head')
    finally:
        with psycopg.connect(ADMIN_URL, autocommit=True) as admin:
            admin.execute(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)')


@pytest.mark.parametrize('text', [
    '李经理表示同意。', '请欧阳总监确认。', '王总经理表示同意。', '马上通知对方。', '由施工方负责清理。',
    '费用要求见附件。', '管理要求见第五条。', '车间负责人确认后发货。', '顾问建议修改。', '申请表示通过。',
])
def test_r4_09_acceptance_negative(text):
    """R4-09 验收（Q-163）：只有姓的称呼与常见词不列为人名。"""
    assert not [s for s in find_suspects(text, []) if s.reason == 'person_name']


@pytest.mark.parametrize(('text', 'name'), [
    ('王博表示同意。', '王博'), ('客户钱虚鸣来电询问交期。', '钱虚鸣'), ('司马虚远先生已签收。', '司马虚远'),
])
def test_r4_09_acceptance_positive(text, name):
    assert [(text[s.start:s.end], s.reason) for s in find_suspects(text, [])] == [(name, 'person_name')]
